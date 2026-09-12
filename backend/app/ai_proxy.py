import asyncio
import base64
import io
import json
import math
import re
import subprocess
import sys
import uuid

import cv2
import httpx
import numpy as np
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from PIL import Image as PILImage
from sqlalchemy.orm import Session

from app import auth, billing, crud, models
from app import content_research
from app.config import (
    BASE_DIR,
    DEV_UNRESTRICTED,
    OPENLUX_API_KEY,
    OPENLUX_BASE_URL,
    VIDU_API_KEY,
    VIDU_BASE_URL,
)
from app.database import get_db
from app.design_tokens import COMPONENT_SIZE
from app.text_metrics import chars_per_line, estimate_text_height, estimate_text_lines
from app import deck_gen, layout_presets
from app.schemas import (
    ChatCompletionRequest,
    ContentResearchRequest,
    DeckPptxRequest,
    DeckRequest,
    DesignElementRequest,
    DesignGenerateRequest,
    DesignLayoutRequest,
    HandoutRequest,
    LayoutPresetRequest,
    ImageGenerationRequest,
    VideoGenerateRequest,
)
from app import handout_categories

router = APIRouter(prefix="/api/ai", tags=["ai-proxy"])

# "参考图生成"产出的背景图自动存一份到这里，供"素材"面板浏览/删除——只存文件名不存完整 URL，
# 换域名/端口不用改数据；实际对外访问路径由 main.py 把这个目录挂到 /api/ai/generated 静态托管。
GENERATED_ASSETS_DIR = BASE_DIR / "generated_assets"
GENERATED_ASSETS_DIR.mkdir(exist_ok=True)


def _persist_asset_bytes(db: Session, user_id: str, category: str, image_bytes: bytes) -> models.GeneratedAsset:
    file_name = f"{uuid.uuid4().hex}.png"
    (GENERATED_ASSETS_DIR / file_name).write_bytes(image_bytes)
    return crud.create_generated_asset(db, user_id, category, file_name)


async def _save_generated_asset(db: Session, user_id: str, category: str, src: str) -> models.GeneratedAsset:
    """src 可能是 data URI（gpt-image-2 直接返回 b64_json 时）或外部 URL（返回 url 时）——
    两种都落盘成本地文件，不满足于存一个可能过期/限流的外链，"自动保存"才是真的能一直看到。"""
    if src.startswith("data:"):
        _, b64data = src.split(",", 1)
        image_bytes = base64.b64decode(b64data)
    else:
        async with httpx.AsyncClient(timeout=30) as client:
            res = await client.get(src)
            res.raise_for_status()
            image_bytes = res.content
    return _persist_asset_bytes(db, user_id, category, image_bytes)


async def _extract_openai_image_bytes(gen_json: dict, error_prefix: str) -> bytes:
    data = (gen_json.get("data") or [None])[0]
    if not data:
        raise HTTPException(status_code=502, detail=f"{error_prefix}未返回图片数据")
    if data.get("b64_json"):
        return base64.b64decode(data["b64_json"])
    url = data.get("url")
    if not url:
        raise HTTPException(status_code=502, detail=f"{error_prefix}未返回可用图片数据")
    async with httpx.AsyncClient(timeout=30) as client:
        res = await client.get(url)
        res.raise_for_status()
        return res.content

# 跟 imageApi.ts 里的 ASPECT_SIZE 保持一致的 5 档比例，图片槽按宽高比就近取一档
_ASPECT_SIZES: list[tuple[str, int, int]] = [
    ("1:1", 768, 768),
    ("3:4", 768, 1024),
    ("4:3", 1024, 768),
    ("9:16", 576, 1024),
    ("16:9", 1024, 576),
]
_DESIGN_MAX_ELEMENTS = 12
# /design/layout 的元素上限比 /design/generate 高不少——那边是"AI 自己编内容"，
# 丢一个装饰性元素问题不大；这边是"用户原文一个字都不能丢"，同一份内容切分成的
# text/icon-list 项数天然就比一版凭空创作的海报多，加上 AI 还会自己加几个装饰性
# 色块/图片，12 个很容易在内容还没排完之前就被截断（实测过：真的截断过一次）。
_LAYOUT_MAX_ELEMENTS = 24


def _require_openlux():
    if not OPENLUX_API_KEY:
        raise HTTPException(status_code=503, detail="AI 服务未配置，请联系管理员")


def _require_vidu():
    if not VIDU_API_KEY:
        raise HTTPException(status_code=503, detail="AI 视频服务未配置，请联系管理员")


async def _post_openlux(url: str, timeout: float, max_retries: int = 2, **kwargs) -> httpx.Response:
    """统一处理超时异常 + 429（"上游负载已饱和"，官方原话是临时性的、稍后重试即可）自动退避重试，
    避免每个转发接口都重复写这段逻辑。detail 里不带中文前缀——前端 authPostJson 自己会按各自场景拼
    "生成接口请求失败"/"消除接口请求失败"这类 label，后端再拼一遍会导致文案重复两遍。"""
    net_err: Exception | None = None
    for attempt in range(max_retries + 1):
        try:
            async with httpx.AsyncClient(timeout=timeout) as client:
                res = await client.post(url, **kwargs)
        except httpx.HTTPError as exc:
            # 网络抖动/超时也重试一次（openlux 偶发；不重试的话用户端一次 504 就卡死）
            net_err = exc
            if attempt < max_retries:
                await asyncio.sleep(1)
                continue
            raise HTTPException(status_code=504, detail=f"请求超时或网络异常，请重试：{exc}")
        if res.status_code != 429 or attempt == max_retries:
            return res
        await asyncio.sleep(3 * (attempt + 1))
    raise HTTPException(status_code=504, detail=f"请求超时或网络异常，请重试：{net_err}")


# 生图/写作/AI消除三个"用户自由输入直达生成能力"的入口专用合规检查——design/generate、
# reference-to-background 这类内部业务逻辑自己拼系统 prompt 调 openlux，不走这层，
# 不需要也不应该被这个检查拦（用户填不了自由文本，风险面完全不同）。
# 不是接的专业内容审核 API（没有这类订阅），是借同一个对话模型做 best-effort 分类——
# 不追求完美，重点是明确挡住"帮忙P证件/仿造官方文件"这类一眼能判断出来的高风险请求。
_ILLEGAL_TEXT_REQUEST_INSTRUCTION = """You are a content-safety classifier for an ordinary consumer design/writing tool (posters, ads, articles, translations, illustrations, etc.). Given the user's request below, decide whether fulfilling it would likely facilitate illegal activity under Chinese law — for example: forging or altering an official ID card/passport/driver's license/household register/certificate/diploma/official seal/government document; counterfeiting currency, stamps, or financial instruments; producing fraudulent contracts, invoices, or medical/legal documents; or otherwise clearly facilitating fraud, forgery, or another crime.

The overwhelming majority of requests to this tool are completely legitimate. Only answer REJECT when illegal intent is reasonably clear from the request itself — when in doubt, or for an ordinary creative/business request, answer ALLOW.

User request:
\"\"\"
{content}
\"\"\"

Respond with EXACTLY one line, no other text:
ALLOW
or
REJECT: <one short Chinese sentence explaining why>"""

_SENSITIVE_DOCUMENT_IMAGE_CHECK_INSTRUCTION = """Look at this image carefully. Does it show — fully, partially, or as a photo/scan/screenshot of one — any of the following:
- a government-issued ID card, passport, driver's license, household register, certificate, diploma, official seal/stamp, or banknote; OR
- a financial/business document such as an invoice (发票), receipt (收据), expense reimbursement form (报销单), bank statement, payslip, or contract?

Respond with EXACTLY one word, no other text: YES or NO."""


async def _moderate_text(content: str) -> None:
    """涉及生图/写作的用户自由文本先过一遍合规判断，命中 REJECT 直接 403。
    分类调用本身失败（超时/网络异常/解析失败）一律当成检查失败拒绝这次请求，
    不静默放行——宁可用户重试一次，不能让检查失效变成事实上没有这层拦截。

    DEV_UNRESTRICTED=1 时整个跳过（仅限本地自测，见 config.py 注释）。"""
    if not content or not content.strip():
        return
    if DEV_UNRESTRICTED:
        print("[DEV_UNRESTRICTED] 跳过文字意图审核", file=sys.stderr, flush=True)
        return
    try:
        res = await _post_openlux(
            f"{OPENLUX_BASE_URL}/chat/completions",
            timeout=30,
            headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
            json={
                "model": "gemini-3-flash-preview",
                "messages": [{"role": "user", "content": _ILLEGAL_TEXT_REQUEST_INSTRUCTION.format(content=content[:2000])}],
            },
        )
        if res.status_code >= 400:
            raise ValueError(f"{res.status_code} {res.text}")
        verdict = res.json()["choices"][0]["message"]["content"].strip()
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=503, detail="内容安全检查失败，请重试")
    if verdict.upper().startswith("REJECT"):
        reason = verdict.split(":", 1)[1].strip() if ":" in verdict else "该请求可能涉及违法内容"
        raise HTTPException(status_code=403, detail=f"请求已被拒绝：{reason}")


async def _check_not_sensitive_document(image_bytes: bytes, media_type: str, feature_label: str) -> None:
    """凡是接受用户上传图片的功能都要过这一层——先判断上传图是不是身份证/证件/公文，或者
    发票/报销单/银行流水这类财务票据，防止被用来篡改/伪造这些材料。所有工具统一走这一个
    分类器，不是只有 AI 消除才防，跟 _moderate_text 一样：分类调用本身失败就拒绝这次请求，
    不静默放行。feature_label 只用来给拒绝提示换个功能名字，检测逻辑完全一样。

    DEV_UNRESTRICTED=1 时整个跳过（仅限本地自测，见 config.py 注释）。"""
    if DEV_UNRESTRICTED:
        print(f"[DEV_UNRESTRICTED] 跳过敏感文件检查：{feature_label}", file=sys.stderr, flush=True)
        return
    # 分类器只需看清"这是不是证件/票据"，缩到 768px 转 JPEG——原图直传大手抄报能到几 MB，
    # base64 塞进 JSON body 上传慢，弱网下 30s 超时 → 504，每个带图的工具全挂
    check_bytes, check_type = image_bytes, media_type
    try:
        im = PILImage.open(io.BytesIO(image_bytes))
        im = im.convert("RGB")
        if max(im.size) > 768:
            r = 768 / max(im.size)
            im = im.resize((max(1, round(im.width * r)), max(1, round(im.height * r))), PILImage.LANCZOS)
        _b = io.BytesIO()
        im.save(_b, "JPEG", quality=78)
        check_bytes, check_type = _b.getvalue(), "image/jpeg"
    except Exception:
        pass
    b64 = base64.b64encode(check_bytes).decode()
    try:
        res = await _post_openlux(
            f"{OPENLUX_BASE_URL}/chat/completions",
            timeout=45,
            headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
            json={
                "model": "gemini-3-flash-preview",
                "messages": [
                    {
                        "role": "user",
                        "content": [
                            {"type": "text", "text": _SENSITIVE_DOCUMENT_IMAGE_CHECK_INSTRUCTION},
                            {"type": "image_url", "image_url": {"url": f"data:{check_type};base64,{b64}"}},
                        ],
                    }
                ],
            },
        )
        if res.status_code >= 400:
            raise ValueError(f"{res.status_code} {res.text}")
        verdict = res.json()["choices"][0]["message"]["content"].strip().upper()
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=503, detail="内容安全检查失败，请重试")
    if verdict.startswith("YES"):
        raise HTTPException(
            status_code=403,
            detail=f"为防止被用于伪造证件/票据，{feature_label}功能不支持处理身份证、证件、发票、报销单等敏感文件类图片",
        )


@router.post("/images/generations")
async def images_generations(
    payload: ImageGenerationRequest,
    user: models.User = Depends(auth.get_current_user),
    db: Session = Depends(get_db),
):
    _require_openlux()
    await _moderate_text(payload.prompt)
    ticket = billing.consume(db, user, "AI生图")
    try:
        res = await _post_openlux(
            f"{OPENLUX_BASE_URL}/images/generations",
            timeout=170,
            headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
            json=payload.model_dump(),
        )
        if res.status_code >= 400:
            raise HTTPException(status_code=502, detail=f"{res.status_code} {res.text}")
        return res.json()
    except Exception:
        billing.refund_ticket(db, user, ticket)
        raise


@router.post("/images/edits")
async def images_edits(
    image: UploadFile = File(...),
    mask: UploadFile = File(...),
    prompt: str = Form(...),
    model: str = Form("gpt-image-2"),
    n: int = Form(1),
    _user: models.User = Depends(auth.get_current_user),
):
    _require_openlux()
    image_bytes = await image.read()
    mask_bytes = await mask.read()
    await _check_not_sensitive_document(image_bytes, image.content_type or "image/png", "AI 消除/文字替换")
    await _moderate_text(prompt)
    res = await _post_openlux(
        f"{OPENLUX_BASE_URL}/images/edits",
        timeout=170,
        headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
        data={"model": model, "prompt": prompt, "n": str(n)},
        files={
            "image": (image.filename or "image.png", image_bytes, image.content_type or "image/png"),
            "mask": (mask.filename or "mask.png", mask_bytes, mask.content_type or "image/png"),
        },
    )
    if res.status_code >= 400:
        raise HTTPException(status_code=502, detail=f"{res.status_code} {res.text}")
    return res.json()


@router.post("/table-to-xlsx")
async def table_to_xlsx(
    image: UploadFile = File(...),
    paper: str = Form("A4"),
    orientation: str = Form("auto"),
    _user: models.User = Depends(auth.get_current_user),
):
    """表格照片 → Excel。让视觉模型把表格读成 JSON 二维数组，openpyxl 生成 xlsx。
    跟 OCR 一样是纯读取/转录，不接敏感文件检测。"""
    from fastapi.responses import StreamingResponse

    from app import table_extract

    _require_openlux()
    image_bytes = await image.read()
    b64 = base64.b64encode(image_bytes).decode()
    media = image.content_type or "image/png"
    res = await _post_openlux(
        f"{OPENLUX_BASE_URL}/chat/completions",
        timeout=120,
        headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
        json={
            "model": "gemini-3-flash-preview",
            "messages": [
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": table_extract.TABLE_INSTRUCTION},
                        {"type": "image_url", "image_url": {"url": f"data:{media};base64,{b64}"}},
                    ],
                }
            ],
        },
    )
    if res.status_code >= 400:
        raise HTTPException(status_code=502, detail=f"表格识别失败：{res.status_code} {res.text[:300]}")
    try:
        content = res.json()["choices"][0]["message"]["content"]
    except Exception:
        raise HTTPException(status_code=502, detail=f"识别服务返回异常：{res.text[:300]}")
    try:
        spec = table_extract.parse_spec(content)
        xlsx_bytes = table_extract.build_xlsx(
            spec,
            paper=paper if paper in ("A4", "A3") else "A4",
            orientation=orientation if orientation in ("auto", "portrait", "landscape") else "auto",
        )
    except ValueError as e:
        raise HTTPException(status_code=422, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"生成 Excel 失败：{type(e).__name__}: {e}")
    return StreamingResponse(
        io.BytesIO(xlsx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=table.xlsx"},
    )


@router.post("/remove-repeated-watermark")
async def remove_repeated_watermark(
    image: UploadFile = File(...),
    box: str = Form(...),  # "x,y,w,h" 相对整图 0~1
    threshold: float = Form(0.45),
    feather: int = Form(3),
    _user: models.User = Depends(auth.get_current_user),
):
    """框选一个水印实例 → 模板匹配找出所有相同的 → 合成蒙版 → inpaint 补掉。
    去水印属于"修改"（不是 OCR 那种纯读取），保留敏感文件检测——不让拿来抹掉证件/执照上的
    防伪水印；检测要调 openlux，所以这个接口也要求 key 配好。实际去水印是纯本地 OpenCV。"""
    from app import watermark_batch

    _require_openlux()
    image_bytes = await image.read()
    await _check_not_sensitive_document(image_bytes, image.content_type or "image/png", "批量去水印")
    try:
        parts = [float(v) for v in box.split(",")]
        assert len(parts) == 4
    except Exception:
        raise HTTPException(status_code=400, detail="框选参数格式不对")
    try:
        out_bytes, count = await asyncio.to_thread(
            watermark_batch.remove_repeated, image_bytes, tuple(parts), threshold, feather
        )
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    b64 = base64.b64encode(out_bytes).decode()
    return {"data": [{"b64_json": b64}], "count": count}


# 参考图→整图背景生成这条链路专用：只要求视觉模型输出"氛围/色调 + 元素类别 + 构图留白说明"
# 这个粒度的风格描述，不要具体坐标/精确外形/可读文字——这是守住"参考图边界判断指南"里
# 版权红线的关键机制，格式严格照抄党建展板那次人工写的 prompt 验证过管用的结构。
_REFERENCE_STYLE_PROMPT_INSTRUCTION = """You are extracting a STYLE-DESCRIPTION prompt from a reference poster image, to be used as input for a text-to-image generation model. Write ONE paragraph in Chinese, following exactly this structure: [整体氛围/光影/色调] + [核心视觉元素类别，用类别名词，不是精确外形] + [构图与留白说明，供后续叠加文字用]。

Strict rules:
- Do NOT include precise coordinates, exact positions, exact sizes/proportions, or counts.
- Do NOT transcribe or describe any specific readable text/words visible in the image.
- Do NOT describe exact rendering details that would let someone reconstruct the image precisely. Describe visual elements only as CATEGORY/TYPE the way a mood-board brief would, not as a blueprint.
- Match the length, tone, and format of this reference example exactly:
"国风，柔和金色光影，飘扬的五星红旗元素、华表、长城剪影、和平鸽、祥云纹样、飘带，简约肌理底纹，画面上方三分之一留白，不能出现文字"
- Output ONLY the style-description paragraph itself, in Chinese, no markdown, no extra commentary, no quotation marks around it.

After that paragraph, on a NEW final line, classify ONLY the RENDERING TECHNIQUE CATEGORY of the reference poster's main title text (never its exact wording, font, or precise shape) using this exact format:
TITLE_STYLE: <effect> <shape>
where <effect> is exactly one of: outline | emboss | neon | plain
and <shape> is exactly one of: straight | arc | wave | ribbon | circle
Pick the closest category by technique only, e.g. a poster with a raised/3D-looking title with soft drop shadow -> "emboss straight"; a title curved along a banner -> "outline arc". If unsure, use "plain straight"."""


_TITLE_EFFECT_MAP = {"outline": "outline", "emboss": "emboss", "neon": "neon", "plain": "none"}
_TITLE_SHAPE_MAP = {"straight": "none", "arc": "arc-up", "wave": "wave", "ribbon": "flag", "circle": "ring"}
_DEFAULT_TITLE_STYLE = {"effect": "none", "warp": "none"}


def _parse_title_style(raw_content: str) -> tuple[str, dict]:
    """从视觉模型输出里拆出末尾的 TITLE_STYLE 分类行，映射成编辑器已有的文字特效/变形预设名。
    只学"手法类别"（描边/浮雕/霓虹 + 直线/拱形/波浪/旗帜/圆环），不涉及具体字形/字体，
    是 [[feedback_reference_image_boundary]] 里"工艺手法可学习复用，具体表达要自己重新实现"
    这条原则在标题文字上的落地。格式解析失败一律回退成"无特效"，不能让格式问题打断生成链路。"""
    lines = raw_content.strip().splitlines()
    if not lines or not lines[-1].strip().upper().startswith("TITLE_STYLE:"):
        return raw_content.strip(), dict(_DEFAULT_TITLE_STYLE)
    style_description = "\n".join(lines[:-1]).strip()
    tokens = lines[-1].split(":", 1)[1].strip().lower().split()
    effect = _TITLE_EFFECT_MAP.get(tokens[0], "none") if tokens else "none"
    warp = _TITLE_SHAPE_MAP.get(tokens[1], "none") if len(tokens) >= 2 else "none"
    return style_description or raw_content.strip(), {"effect": effect, "warp": warp}


@router.post("/design/reference-to-background")
async def design_reference_to_background(
    image: UploadFile = File(...),
    user: models.User = Depends(auth.get_current_user),
    db: Session = Depends(get_db),
):
    """参考图 → 风格描述（视觉模型，只提取氛围/元素类别/构图留白，不提取可判定为复刻的精确细节）
    → 整图背景生成（gpt-image-2）。三次跨风格实测（国风红金/卡通插画/极简科技）验证过，
    质量稳定在 tpl-board-party-building 那次人工调用的水准，全部一次生成成功。
    标题/副标题文字层不在这里加——那部分复用前端已有的文字编辑能力，用户自己调。"""
    _require_openlux()
    image_bytes = await image.read()
    media_type = image.content_type or "image/png"
    await _check_not_sensitive_document(image_bytes, media_type, "参考图生成")
    b64_in = base64.b64encode(image_bytes).decode()
    ticket = billing.consume(db, user, "参考图生成")

    try:
        return await _do_reference_to_background(db, user, media_type, b64_in)
    except Exception:
        billing.refund_ticket(db, user, ticket)
        raise


async def _do_reference_to_background(db, user, media_type: str, b64_in: str):
    vision_res = await _post_openlux(
        f"{OPENLUX_BASE_URL}/chat/completions",
        timeout=60,
        headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
        json={
            "model": "gemini-3-flash-preview",
            "messages": [
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": _REFERENCE_STYLE_PROMPT_INSTRUCTION},
                        {"type": "image_url", "image_url": {"url": f"data:{media_type};base64,{b64_in}"}},
                    ],
                }
            ],
        },
    )
    if vision_res.status_code >= 400:
        raise HTTPException(status_code=502, detail=f"参考图风格分析失败：{vision_res.status_code} {vision_res.text}")
    raw_content = vision_res.json()["choices"][0]["message"]["content"].strip()
    style_description, title_style = _parse_title_style(raw_content)

    gen_res = await _post_openlux(
        f"{OPENLUX_BASE_URL}/images/generations",
        timeout=170,
        headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
        json={"model": "gpt-image-2", "prompt": style_description, "n": 1, "size": "1024x1536"},
    )
    if gen_res.status_code >= 400:
        raise HTTPException(status_code=502, detail=f"背景图生成失败：{gen_res.status_code} {gen_res.text}")
    data = (gen_res.json().get("data") or [None])[0]
    if not data:
        raise HTTPException(status_code=502, detail="背景图生成未返回图片数据")
    src = data.get("url") or (f"data:image/png;base64,{data['b64_json']}" if data.get("b64_json") else None)
    if not src:
        raise HTTPException(status_code=502, detail="背景图生成未返回可用图片数据")

    asset_id = None
    try:
        asset = await _save_generated_asset(db, user.id, "reference-background", src)
        asset_id = asset.id
    except Exception:
        pass  # 自动保存失败不能拖累主流程——用户还是要拿到刚生成的背景图，大不了这次没存进素材库

    return {"backgroundSrc": src, "styleDescription": style_description, "titleStyle": title_style, "assetId": asset_id}


@router.post("/design/handout")
async def design_handout(
    payload: HandoutRequest,
    user: models.User = Depends(auth.get_current_user),
    db: Session = Depends(get_db),
):
    """手抄报/黑板报一键生成。很多用户（尤其带娃的家长）不会写 prompt，所以不让 AI 自由发挥
    整个结构——先选好分类（app/handout_categories.py，每类自带一套约定俗成的板块标题+配色+
    边框风格），AI 只负责往已经定好的框架里填具体内容，更像"组装"不是"写作文"，出错自由度小。
    AI 只画左半边的主体插画（右半 + 顶部留白），正文和艺术大标题是叠加的独立文字图层，
    不烧进图片里，改文字/重排版不影响插画。同时用 OpenCV 从彩色版提一张黑白线稿版，
    家长可以照着彩色版给孩子涂色。with_content=False 出纯涂色版（只有插画+标题）。"""
    _require_openlux()
    cat = handout_categories.get_category(payload.category)
    style = handout_categories.get_style(payload.style)
    border = handout_categories.get_border(payload.border)
    topic = payload.topic.strip()
    custom = payload.custom_prompt.strip()[:500]
    if custom:
        await _moderate_text(custom)
    W, H, wc = payload.canvas_width, payload.canvas_height, payload.with_content

    if payload.layered:
        # 可拆分版要跑十来次生图，3~5 分钟——同步等会撞 nginx/网关超时，改成"下单 → 轮询"。
        ticket = billing.consume(db, user, "手抄报拆分")
        job_id = uuid.uuid4().hex
        _HANDOUT_JOBS[job_id] = {"status": "pending", "user_id": user.id}
        _prune_handout_jobs()
        asyncio.create_task(
            _run_handout_layered_job(job_id, user.id, ticket, cat, style, border, topic, W, H, wc, custom)
        )
        return {"jobId": job_id}

    ticket = billing.consume(db, user, "手抄报生成")
    try:
        return await _do_handout(db, user, cat, style, border, topic, W, H, wc, custom)
    except Exception:
        billing.refund_ticket(db, user, ticket)
        raise


# job_id -> {status: pending|done|error, result?, detail?, user_id, ts}
_HANDOUT_JOBS: dict[str, dict] = {}


def _prune_handout_jobs(keep: int = 40) -> None:
    if len(_HANDOUT_JOBS) <= keep:
        return
    done = [k for k, v in _HANDOUT_JOBS.items() if v.get("status") != "pending"]
    for k in done[: max(0, len(_HANDOUT_JOBS) - keep)]:
        _HANDOUT_JOBS.pop(k, None)


async def _run_handout_layered_job(job_id, user_id, ticket, cat, style, border, topic, W, H, wc, custom=""):
    """后台跑「可拆分手抄报」，结果塞进 _HANDOUT_JOBS。自带 db session（不能用请求的）。"""
    from app.database import SessionLocal

    db = SessionLocal()
    try:
        user = db.query(models.User).filter(models.User.id == user_id).first()
        if user is None:
            raise RuntimeError("用户不存在")
        result = await _do_handout_layered(db, user, cat, style, border, topic, W, H, wc, custom)
        _HANDOUT_JOBS[job_id] = {"status": "done", "result": result, "user_id": user_id}
    except HTTPException as e:
        try:
            user = db.query(models.User).filter(models.User.id == user_id).first()
            if user:
                billing.refund_ticket(db, user, ticket)
        except Exception:
            pass
        _HANDOUT_JOBS[job_id] = {"status": "error", "detail": str(e.detail), "user_id": user_id}
    except Exception as e:  # noqa: BLE001
        try:
            user = db.query(models.User).filter(models.User.id == user_id).first()
            if user:
                billing.refund_ticket(db, user, ticket)
        except Exception:
            pass
        _HANDOUT_JOBS[job_id] = {"status": "error", "detail": f"生成失败：{e}", "user_id": user_id}
    finally:
        db.close()


@router.get("/design/handout/job/{job_id}")
async def get_handout_job(job_id: str, user: models.User = Depends(auth.get_current_user)):
    job = _HANDOUT_JOBS.get(job_id)
    if not job or job.get("user_id") != user.id:
        raise HTTPException(status_code=404, detail="任务不存在或已过期")
    if job["status"] == "pending":
        return {"status": "pending"}
    if job["status"] == "error":
        return {"status": "error", "detail": job.get("detail", "生成失败")}
    return {"status": "done", "result": job["result"]}


def _to_lineart(image_bytes: bytes) -> bytes:
    """从彩色版提取黑白线稿——家长打印出来给孩子照着涂色。用自适应阈值而不是 Canny：
    Canny 出来是断线的细描边，adaptiveThreshold 出来是连续闭合的粗轮廓，更像儿童填色书。"""
    arr = np.frombuffer(image_bytes, np.uint8)
    img = cv2.imdecode(arr, cv2.IMREAD_COLOR)
    if img is None:
        raise ValueError("decode failed")
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    gray = cv2.medianBlur(gray, 3)
    edges = cv2.adaptiveThreshold(
        gray, 255, cv2.ADAPTIVE_THRESH_MEAN_C, cv2.THRESH_BINARY, 9, 5
    )
    ok, buf = cv2.imencode(".png", edges)
    if not ok:
        raise ValueError("encode failed")
    return buf.tobytes()


# ─────────────────────────────────────────────────────────────────────────────
# 手抄报「可拆分元素版」：AI 出整张白底手抄报 → 视觉模型框出每个图画元素 →
# 逐个裁下来、把连着边界的白底 flood-fill 成透明 → 铺成一堆可单独拖动/替换的图层。
# ─────────────────────────────────────────────────────────────────────────────
def _cutout_white_bg(image_bytes: bytes, white_thresh: int = 236) -> bytes:
    """把一张白底小图的背景抠成透明。从四条边界向内 flood-fill 连通的近白色区域——
    只吃"跟边框连着的背景白"，物体内部的白（眼白/高光）不受影响。"""
    arr = np.frombuffer(image_bytes, np.uint8)
    img = cv2.imdecode(arr, cv2.IMREAD_COLOR)
    if img is None:
        raise ValueError("decode failed")
    h, w = img.shape[:2]
    mask = np.zeros((h + 2, w + 2), np.uint8)
    tol = max(6, 255 - white_thresh)
    flags = 8 | cv2.FLOODFILL_MASK_ONLY | (255 << 8)
    step_x = max(1, w // 60)
    step_y = max(1, h // 60)
    seeds = (
        [(x, 0) for x in range(0, w, step_x)]
        + [(x, h - 1) for x in range(0, w, step_x)]
        + [(0, y) for y in range(0, h, step_y)]
        + [(w - 1, y) for y in range(0, h, step_y)]
    )
    for sx, sy in seeds:
        if mask[sy + 1, sx + 1] == 0 and int(img[sy, sx].min()) >= white_thresh:
            cv2.floodFill(img, mask, (sx, sy), 0, (tol, tol, tol), (tol, tol, tol), flags)
    bg = mask[1:-1, 1:-1]
    alpha = np.where(bg > 0, 0, 255).astype(np.uint8)
    # 往里收 1px 去白边，再轻羽化
    alpha = cv2.erode(alpha, np.ones((2, 2), np.uint8), iterations=1)
    alpha = cv2.GaussianBlur(alpha, (3, 3), 0)
    b, g, r = cv2.split(img)
    ok, buf = cv2.imencode(".png", cv2.merge([b, g, r, alpha]))
    if not ok:
        raise ValueError("encode failed")
    return buf.tobytes()


async def _gen_image_bytes(prompt: str, size: str, *, attempts: int = 2, timeout: int = 160) -> bytes:
    """gpt-image-2 生成一张图，返回原始字节；重试 attempts 次都失败抛 502。"""
    last = ""
    for _ in range(attempts):
        try:
            res = await _post_openlux(
                f"{OPENLUX_BASE_URL}/images/generations",
                timeout=timeout,
                headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
                json={"model": "gpt-image-2", "prompt": prompt, "n": 1, "size": size},
            )
            if res.status_code >= 400:
                last = f"{res.status_code} {res.text[:200]}"
                continue
            return await _extract_openai_image_bytes(res.json(), "图片")
        except HTTPException as e:
            last = str(e.detail)
        except Exception as e:  # noqa: BLE001
            last = str(e)
    raise HTTPException(status_code=502, detail=f"图片生成失败：{last}")


_GEMINI_IMAGE_MODEL = "gemini-3.1-flash-image"


async def _gemini_image(prompt: str, ref_b64: str | None = None, *, attempts: int = 3, timeout: int = 150) -> bytes:
    """gemini 图像模型（对话式生成/编辑，原生输出透明 PNG）。ref_b64 传参考图。
    429「上游拥挤」会退避重试；全失败抛 502。"""
    content: list[dict] = [{"type": "text", "text": prompt}]
    if ref_b64:
        content.append({"type": "image_url", "image_url": {"url": f"data:image/png;base64,{ref_b64}"}})
    last = ""
    for i in range(attempts):
        try:
            res = await _post_openlux(
                f"{OPENLUX_BASE_URL}/chat/completions",
                timeout=timeout,
                headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
                json={"model": _GEMINI_IMAGE_MODEL, "messages": [{"role": "user", "content": content}]},
            )
            if res.status_code == 429 or res.status_code >= 500:
                last = f"{res.status_code}"
                await asyncio.sleep(2 + i * 3)
                continue
            if res.status_code >= 400:
                raise HTTPException(status_code=502, detail=f"元素生成失败：{res.status_code} {res.text[:160]}")
            c = res.json().get("choices", [{}])[0].get("message", {}).get("content") or ""
            m = re.search(r"base64,([A-Za-z0-9+/=]+)", c)
            if not m:
                last = "no image in response"
                await asyncio.sleep(1)
                continue
            return base64.b64decode(m.group(1))
        except HTTPException:
            raise
        except Exception as e:  # noqa: BLE001
            last = str(e)
            await asyncio.sleep(1)
    raise HTTPException(status_code=502, detail=f"元素生成失败：{last}")


@router.post("/design/element")
async def design_element(
    payload: DesignElementRequest,
    user: models.User = Depends(auth.get_current_user),
    db: Session = Depends(get_db),
):
    """手抄报可拆分元素版——替换单个元素：一句提示词 → 一张透明底小图。"""
    _require_openlux()
    prompt = payload.prompt.strip()
    if not prompt:
        raise HTTPException(status_code=400, detail="请描述要生成的元素")
    style = handout_categories.get_style(payload.style)
    ticket = billing.consume(db, user, "素材生成")
    try:
        gen_prompt = (
            f"画一个「{prompt}」，{style['prompt']}，粗黑描边，只有这一个物体、完整居中、占满画面，"
            "纯透明背景，不要文字、不要边框、不要地面阴影、不要多个物体。"
        )
        raw = await _gemini_image(gen_prompt, attempts=3, timeout=120)
        asset = _persist_asset_bytes(db, user.id, "handout-element", raw)
        return {"src": f"/api/ai/generated/{asset.file_name}", "assetId": asset.id}
    except Exception:
        billing.refund_ticket(db, user, ticket)
        raise


_LINEART_PROMPT = (
    "把这张图转换成干净的黑白线稿涂色页。要求："
    "只保留清晰、连贯、粗细均匀的纯黑色描边线条；"
    "去掉所有颜色、灰阶、阴影、纹理和网点底纹；"
    "同一条轮廓只画一条线，不要出现描边两侧各一条的双线；"
    "图中的文字也要转成清晰、笔画完整、可正常阅读的黑色线条，不要糊成一团；"
    "纯白色背景；构图、每个元素的位置和大小跟原图保持一致，不要新增或删减内容。"
)


@router.post("/design/lineart")
async def design_lineart(
    image: UploadFile = File(...),
    user: models.User = Depends(auth.get_current_user),
    db: Session = Depends(get_db),
):
    """任意彩色图 → 干净黑白线稿（AI 版，效果接近豆包）。本地 canvas 那版是自适应阈值，
    对蜡笔/水彩纹理和文字处理不好；这里用 gemini 图像模型重画成线稿。前端再按「深浅」
    把黑线压成任意浅度铺到透明底上。"""
    _require_openlux()
    image_bytes = await image.read()
    await _check_not_sensitive_document(image_bytes, image.content_type or "image/png", "转线稿")
    ticket = billing.consume(db, user, "AI线稿")
    try:
        b64 = base64.b64encode(image_bytes).decode()
        raw = await _gemini_image(_LINEART_PROMPT, ref_b64=b64, attempts=3, timeout=120)
        asset = _persist_asset_bytes(db, user.id, "lineart", raw)
        return {"src": f"/api/ai/generated/{asset.file_name}", "assetId": asset.id}
    except Exception:
        billing.refund_ticket(db, user, ticket)
        raise


# 大纲 JSON 结构 + 版式规则（"填主题"和"传资料"两条链路共用）
_DECK_JSON_SPEC = (
    "只返回一个严格的 JSON 对象，不要 markdown 代码块、不要多余说明，形如：\n"
    '{"title":"演示标题","subtitle":"一句副标题",'
    '"palette":["#主色","#强调色","#主色深","#背景浅色","#正文深灰"],'
    '"mood":"用一句话描述整体视觉基调，例：庄重大气的党政红金风、简洁现代的科技蓝",'
    '"cover_image_prompt":"给封面配一张 16:9 专业 PPT 封面设计图的提示词",'
    '"section_image_prompt":"给章节过渡页配一张 16:9 氛围图的提示词，风格跟封面一致",'
    '"content_image_prompt":"给正文页配一张几乎纯白、只角落有极淡装饰的底图提示词",'
    '"cover_features":[{"value":"4K","label":"超清影像","en":"4K Ultra HD"}],'
    '"photo_prompts":["一张跟主题强相关的写实照片的英文提示词","另一张…"],'
    '"sections":[{"heading":"章节标题","en":"章节英文短标题(全大写,2~4词)",'
    '"slides":[{"layout":"版式类型","title":"小标题","en":"英文短标题(全大写,1~3词)","intro":"1~2句导语,可空","bullets":["要点一","要点二"],"image":0}]}]}\n'
    "\n【关键】每个 slide 必须先判断内容最适合哪种版式,填 layout 字段(只做这道选择题,不要输出坐标/字号)。"
    "可选 layout 及对应要填的数据字段:\n"
    '- "cards"：2~3 个并列要点(最常用)。填 bullets(2~3 条,每条 12~40 字)\n'
    '- "list"：4~6 个要点,逻辑并列或递进。填 bullets(4~6 条)\n'
    '- "quote"：一句核心观点/口号/金句,需要留白强调。填 bullets(正好 1 条)\n'
    '- "timeline"：有时间/阶段/步骤先后顺序。填 bullets(3~5 步,按顺序)\n'
    '- "big_number"：这一页核心就是一个关键数字。填 big_number:{"value":"85%","label":"客户满意度","note":"一句补充说明"}\n'
    '- "stats"：2~4 个并列的关键指标。填 data:{"kind":"stat","items":[{"label":"标签","value":85}]}\n'
    '- "rings"：2~5 个百分比数据,适合做成环形进度。填 data:{"kind":"ring","items":[{"label":"标签","value":75}]}(value 是 0~100 的数)\n'
    '- "bar"：多项数值需要横向对比。填 data:{"kind":"bar","items":[{"label":"标签","value":85}]}(数字要真实,编不出别用)\n'
    '- "line"：3~7 个时间点/阶段的数值走势(涨跌趋势、预测曲线)。填 data:{"kind":"line","items":[{"label":"2024","value":1050}]}(数字要真实或合理推算,不要瞎编精确到个位的假数字)\n'
    '- "table"：适合用表格对比的多行多列结构化信息(比如多个方案/场景在几个维度上的对比)。'
    '填 table:{"columns":["列名1","列名2",...](2~5列),"rows":[["单元格","单元格",...],...]}(3~6 行,每行长度=columns 长度,第一列通常是行标题)\n'
    '- "spoke"：围绕一个核心概念展开 3~6 个方面/维度/组成部分,彼此并列。title 写核心概念,填 bullets(3~6 条,每条 4~14 字的短语)\n'
    '- "hive"：一组 3~6 个并列的能力/模块/要素,想突出"体系感"。title 写体系名,填 bullets(3~6 条,每条 2~8 字)\n'
    '- "cycle"：3~5 个环节首尾相接、循环往复的闭环流程(区别于 timeline 的单向推进)。title 写循环名,填 bullets(3~5 条,按顺序)\n'
    '- "tree"：一个主干话题分出 2~6 个分支方向/子类。填 bullets(2~6 条,每条 2~10 字)\n'
    '- "diamond"：正好 3~4 个并列要点,想要点缀感、不想用卡片。填 bullets(3~4 条,每条 2~10 字)\n'
    '- "bulb"：3~6 个"想法/亮点/启发"类要点,强调创意感(区别于 spoke 的中性并列)。填 bullets(3~6 条,每条 2~10 字)\n'
    '- "compare"：两个对象/方案/时期的对照。填 compare:{"left":{"heading":"左栏标题","points":["要点"]},"right":{"heading":"右栏标题","points":["要点"]}}(每栏 3~4 条)\n'
    '- "matrix"：按两个维度分成四类。填 matrix:{"xLabel":"横轴","yLabel":"纵轴","cells":[{"title":"象限名","items":["要点"]}]}(正好 4 个 cell)\n'
    '- "swot"：专门的 SWOT 态势分析。填 swot:{"s":[],"w":[],"o":[],"t":[]}(每项 2~4 条)\n'
    "cover / section_divider / closing 由系统自动排,不用你选。\n"
    "分布要求:同一份大纲里 layout 至少出现 4 种以上,不要每页都是 cards;"
    "compare/matrix/swot/big_number/spoke/hive/cycle/tree/diamond/bulb/line/table 各最多 1~2 页,只在真契合时用；"
    "table/line 涉及具体数字/结构化对比,内容里有靠谱数据支撑才用,别为了凑版式种类编数字。\n"
    "palette 必须是 5 个协调的十六进制色，符合主题气质、对比度足够（正文色要能在背景浅色上看清）；"
    "en 字段是给版式当装饰小字用的英文，要贴切、地道。\n"
    "cover_image_prompt：描述一张能直接当商业 PPT 封面的完整设计图。参考市面成品模板的做法——"
    "一张跟主题强相关的高质量主视觉（产品渲染图 / 行业实景 / 象征元素 / 意境画面）占据画面右侧约 60%，"
    "配简洁的斜切几何色块、光效、细线条装饰；符合 mood 的色调；"
    "关键约束：画面左侧约 40% 留出干净、低细节、纯色或浅色的区域给标题；整张图绝对不要任何文字/字母/数字/logo。\n"
    "section_image_prompt：跟封面同一套视觉语言的章节过渡氛围图，主色调铺底 + 主视觉元素，左下约一半区域留干净。\n"
    "content_image_prompt：正文页底图，跟 cover_image_prompt 是**同一套视觉语言**但淡得多——"
    "融入跟主题相关的淡淡意象（相关器物剪影 / 场景轮廓 / 象征元素）+ 细线条 / 半透明几何点缀，装饰主要放在**画面下方和左右两侧、四角**；"
    "顶部约 15% 和正中横向约 55% 高的区域要明显更淡更干净（放标题和文字卡片）。整体明亮、浅色调、留白足，不要深色背景、不要写实照片、不要浓重色块、不要文字。整篇正文页复用这一张。\n"
    "三个 prompt 都写具体，别堆空泛形容词。\n"
    "cover_features：如果主题是产品 / 方案 / 服务发布，且能提炼出 3~4 个亮点规格，就填这个数组"
    "（每项 value=数字或短词、label=中文说明 4~6 字、en=英文，例 {\"value\":\"46分钟\",\"label\":\"超长续航\",\"en\":\"46 Min Flight\"}）；"
    "不是发布类主题就设为空数组 []。\n"
    "文案排版规范（办公稿标准，务必遵守）：所有中文标点用全角（，。、；：？！“”（）），不要用半角逗号句号；"
    "中文字符之间不加空格；每条 bullet 和 intro 都是完整通顺的句子、以句号结尾；title/heading 是短语、结尾不加标点。\n"
    "photo_prompts：给这份 PPT 配 4~6 张写实照片的英文提示词（会用 AI 生成，嵌在几何图框里）。"
    "每条描述一张跟主题强相关的高质量摄影照片——单一清晰主体（产品实拍 / 行业场景 / 人物工作 / 环境实景），"
    "自然真实的光线，主体居中或偏一侧留出干净空间，画面绝对不要任何文字/水印/logo/拼贴/示意图。"
    "主题实在不适合配实拍照片（纯理论 / 纯数据）就给空数组 []。\n"
    "配图分配：在 2~4 个内容契合的普通 slide（有 bullets 的）上加 \"image\": 照片编号（0 起的整数，对应 photo_prompts 里第几条）。"
    "可以另外挑 1 个 slide 把 layout 设成 \"gallery\" 并加 \"images\": [编号,编号,编号]（正好 3 张，bullets 写这 3 张的短说明）；"
    "也可以挑 1 个 slide 把 layout 设成 \"hive\" 并加 \"images\": [编号,...]（3~6 张，bullets 写每张一句说明），做成蜂窝嵌照片。"
    "一张照片最多用一次；图表页 / 对比页 / SWOT / matrix / big_number 不放图；不契合宁可不放。"
)


_DECK_PHOTO_RULE = (
    "\n用户还上传了以下真实照片（编号从 0 开始）：\n{photo_list}\n"
    "这种情况【忽略 photo_prompts】，photo_prompts 直接给空数组 []；配图只用用户上传的这些照片。"
    "对于内容跟某张照片契合的普通 slide（有 bullets 的），加一个字段 \"image\": 照片编号（整数）。"
    "一张照片最多用在一页；不契合就不要硬配，宁可这页不放图。图表页/对比页/SWOT 页不要放图。"
    "如果有一张照片适合当封面主图，在顶层加 \"cover_image\": 编号。"
)


async def _gen_deck_outline(
    topic: str,
    sections: int,
    extra: str = "",
    material: str = "",
    photo_tags: list[str] | None = None,
    ref_layouts: list[str] | None = None,
    ref_density: str = "",
) -> dict:
    """主题（或整份资料）→ PPT 大纲 JSON（含 title/subtitle/sections + palette + mood）。
    material 非空时走"重组资料"模式：标题/内容全部从资料提炼，不新增资料里没有的信息。
    photo_tags 非空时让模型给合适的 slide 标 "image":编号。
    ref_layouts / ref_density：参考图归类出来的通用版式偏好，只做倾向性引导，内容不契合就不用。"""
    extra_line = f"用户补充要求：{extra}。\n" if extra else ""
    ref_line = ""
    if ref_layouts:
        _dz = {"airy": "整体偏留白、每页别塞太满", "packed": "每页信息量做足、少留白、多用图解页", "balanced": "疏密适中"}
        ref_line = (
            f"【参考风格倾向】用户挑了一套喜欢的模板，它常用这几类通用图解：{('、'.join(ref_layouts))}。"
            f"在内容确实契合时，优先从这几种里选 layout（别硬套，不契合就按内容本身最合适的来）；"
            f"排版{_dz.get(ref_density, '疏密适中')}。这只是风格倾向，版面/尺寸/位置全部由我们自己的引擎决定。\n"
        )
    photo_rule = ""
    if photo_tags:
        photo_list = "\n".join(f"[{i}] {t}" for i, t in enumerate(photo_tags))
        photo_rule = _DECK_PHOTO_RULE.format(photo_list=photo_list)
    if material:
        prompt = (
            "你是资深 PPT 设计师。下面【资料原文】是用户准备好的素材，请把它重组成一份逻辑清晰的幻灯片大纲，"
            "并给出配套视觉方案。\n"
            "硬性要求：标题、章节、要点全部从资料里提炼和归纳，可以精简、改写得更书面、合并同类项，"
            "但不得新增资料里没有的事实、数据或观点，不要脑补。资料里出现的数字/占比要保留并可做成图表页。\n"
            f"章节数：资料结构清晰就按它自然的段落数（2~6 个）来；否则归纳成约 {sections} 个章节。"
            "每个 section 下 2~4 个 slides，普通 slide 3~5 条 bullets、每条 15~45 字。\n"
            f"{extra_line}{ref_line}"
            f"{_DECK_JSON_SPEC}{photo_rule}\n"
            f"【资料原文】\n{material}"
        )
    else:
        prompt = (
            f"你是资深 PPT 设计师。为主题「{topic}」写一份幻灯片大纲，并给出配套的视觉方案。\n"
            f"{extra_line}{ref_line}"
            f"{_DECK_JSON_SPEC}{photo_rule}\n"
            f"要求：sections 生成 {sections} 个；每个 section 下 2~3 个 slides；普通 slide 配 3~5 条 bullets，"
            "每条 20~45 字，具体、准确、书面语，不空话套话；title/heading 精炼；涉及事实或数据要可靠。"
        )
    res = await _post_openlux(
        f"{OPENLUX_BASE_URL}/chat/completions",
        timeout=180,  # 后台 job 里跑，不受 nginx 网关超时限制，给大模型足够时间
        headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
        json={"model": "gemini-3-flash-preview", "messages": [{"role": "user", "content": prompt}]},
    )
    if res.status_code >= 400:
        raise HTTPException(status_code=502, detail=f"大纲生成失败：{res.status_code} {res.text[:160]}")
    raw = res.json().get("choices", [{}])[0].get("message", {}).get("content", "")
    m = re.search(r"\{[\s\S]*\}", raw)
    try:
        data = json.loads(m.group(0) if m else raw)
        assert isinstance(data.get("sections"), list) and data["sections"]
    except Exception:
        raise HTTPException(status_code=502, detail="大纲解析失败，请重试")
    return data


# 封面/章节页整张设计图（对标优品PPT那种成品封面）——图里的安全与画质约束
_DECK_ART_GUARD = (
    "16:9 横版宽屏，专业商业 PPT 的成品设计图，构图讲究、有设计感、画质高清。"
    "主色调：{pal}。{mood}。\n"
    "绝对禁止：任何文字、字母、数字、水印、logo、国徽/党徽/警徽等国家标志、二维码；"
    "不要整圈边框相框；不要低俗或敏感内容。"
)
_DECK_ART_FALLBACK = {
    "cover": (
        "一张大气的商业 PPT 封面：与「{topic}」主题贴切的主视觉放在画面右侧，"
        "配简洁的几何色块和光影装饰；画面左侧到中部约 55% 是干净、低细节的浅色留白区（留给标题）。"
    ),
    "section": (
        "一张与封面统一风格的章节过渡页氛围图：主色铺底，右上方有大的同色系半透明几何装饰，"
        "左下角约一半区域保持干净留白。"
    ),
    # 内容页底图：整篇复用同一张——跟封面同一套视觉语言的极淡版，还要跟主题相关
    "content": (
        "一张浅色调的商业 PPT 正文页背景，跟封面是同一套视觉风格但淡得多：整体明亮的浅灰蓝白色调，"
        "融入跟「{topic}」主题相关的淡淡意象元素（如相关的器物剪影 / 场景轮廓 / 象征符号）+ 细线条 / 半透明几何点缀，"
        "这些装饰主要放在画面下方和左右两侧、四角；"
        "顶部约 15% 和正中间横向约 55% 高的区域要明显更淡、更干净（放标题和文字卡片，不能被压住）；"
        "整体克制、留白充足，不要浓重色块、不要深色背景、不要写实照片、不要文字。"
    ),
}


async def _gen_deck_cover_kit(outline: dict, want_content_bg: bool = True) -> dict:
    """生成封面 + 章节页（整张设计图）+ 内容页底图（极淡、整篇复用）。返回 {kind: bytes}。"""
    pal = "、".join(str(c) for c in (outline.get("palette") or [])[:4]) or "自定协调配色"
    mood = (outline.get("mood") or "简洁现代的商务风").strip()
    topic = (outline.get("title") or "").strip()
    guard = _DECK_ART_GUARD.format(pal=pal, mood=mood)
    plan = {
        "cover": (outline.get("cover_image_prompt") or "").strip()
        or _DECK_ART_FALLBACK["cover"].format(topic=topic or "演示主题"),
        "section": (outline.get("section_image_prompt") or "").strip() or _DECK_ART_FALLBACK["section"],
    }
    if want_content_bg:
        plan["content"] = (outline.get("content_image_prompt") or "").strip() or _DECK_ART_FALLBACK[
            "content"
        ].format(topic=topic or "演示主题")
    kit: dict[str, bytes] = {}
    for kind, body in plan.items():
        try:
            kit[kind] = await _gen_image_bytes(f"{body}\n{guard}", "1536x1024", attempts=2, timeout=150)
        except Exception:
            if kind in ("cover", "section"):
                raise  # 封面/章节图必须有
            # 内容底图失败就算了，用代码画的淡纹
    return kit


def _persist_bg_kit(db: Session, user_id: str, raw_kit: dict) -> dict:
    out = {}
    for kind, raw in raw_kit.items():
        a = _persist_asset_bytes(db, user_id, "deck-bg", raw)
        out[kind] = f"/api/ai/generated/{a.file_name}"
    return out


_BG_DETAIL_SEM = asyncio.Semaphore(4)  # 按章节/按页配图时限流，别一次性把上游打爆


async def _gen_per_page_content_bg(outline: dict, detail: str, db: Session, user_id: str) -> str | None:
    """正文底图不整篇复用一张，改成按章节(detail=section)或按页(detail=slide)各生成一张，
    挂到对应 slide 的 "bg" 字段——templates.ts 里 sl.bg 优先，没有才退回共用那张。
    每张的 prompt = 跟封面同一套视觉语言的基础描述 + 这页/这章节的主题当"呼应"提示，
    保证风格统一、内容各不相同。并发但限流；单张失败跳过，不影响其它页。
    返回第一张生成成功的 URL，当整体兜底（万一某页没配上）。"""
    pal = "、".join(str(c) for c in (outline.get("palette") or [])[:4]) or "自定协调配色"
    mood = (outline.get("mood") or "简洁现代的商务风").strip()
    topic = (outline.get("title") or "").strip()
    guard = _DECK_ART_GUARD.format(pal=pal, mood=mood)
    base = (outline.get("content_image_prompt") or "").strip() or _DECK_ART_FALLBACK["content"].format(
        topic=topic or "演示主题"
    )

    groups: list[tuple[list[dict], str]] = []  # (这组共用一张图的 slide 列表, 主题提示词)
    for sec in outline.get("sections") or []:
        heading = (sec.get("heading") or "").strip()
        slides = [s for s in (sec.get("slides") or []) if isinstance(s, dict)]
        if not slides:
            continue
        if detail == "section":
            groups.append((slides, heading))
        else:  # "slide"
            for sl in slides[:4]:
                hint = (sl.get("title") or heading).strip()
                groups.append(([sl], hint))

    async def gen_one(hint: str) -> bytes | None:
        prompt = f"{base}，画面意象呼应「{hint}」\n{guard}" if hint else f"{base}\n{guard}"
        async with _BG_DETAIL_SEM:
            try:
                return await _gen_image_bytes(prompt, "1536x1024", attempts=1, timeout=150)
            except Exception:
                return None

    raws = await asyncio.gather(*[gen_one(hint) for _, hint in groups])
    fallback_url: str | None = None
    for (slides, _), raw in zip(groups, raws):
        if not raw:
            continue
        asset = _persist_asset_bytes(db, user_id, "deck-bg", raw)
        url = f"/api/ai/generated/{asset.file_name}"
        if fallback_url is None:
            fallback_url = url
        for sl in slides:
            sl["bg"] = url
    return fallback_url


_DECK_PHOTO_GUARD = (
    " Realistic professional photograph, natural lighting, single clear subject, "
    "some clean negative space, no text, no watermark, no logo, no collage, not an illustration."
)


async def _gen_deck_spot_photos(prompts: list[str], db: Session, user_id: str) -> list[dict]:
    """一组英文提示词 → 逐张 AI 生成写实照片，落盘。返回 [{url, tag}]（跟上传照片同结构，
    直接喂 _attach_deck_photos）。单张失败跳过，整体不致命。"""
    out: list[dict] = []
    for p in prompts[:6]:
        p = str(p or "").strip()
        if not p:
            continue
        try:
            raw = await _gen_image_bytes(p + _DECK_PHOTO_GUARD, "1024x1024", attempts=2, timeout=150)
        except Exception:
            continue
        a = _persist_asset_bytes(db, user_id, "deck-photo", raw)
        out.append({"url": f"/api/ai/generated/{a.file_name}", "tag": p[:60]})
    return out


async def _run_deck_job(
    job_id, user_id, ticket, topic, n, theme, extra, ai_bg, material="", photos=None, palette=None,
    ref_layouts=None, ref_density="", ref_motif="", bg_detail="shared",
):
    from app.database import SessionLocal

    photos = photos or []
    db = SessionLocal()
    try:
        user = db.query(models.User).filter(models.User.id == user_id).first()
        if user is None:
            raise RuntimeError("用户不存在")
        outline = await _gen_deck_outline(
            topic, n, extra, material, [p["tag"] for p in photos] if photos else None,
            ref_layouts=ref_layouts, ref_density=ref_density,
        )
        outline = deck_gen.normalize_outline_text(outline)
        outline = deck_gen.apply_theme_palette(outline, theme, palette)
        # 参考图的密度/主装饰形状：只塞进 outline 供前端排版引擎读，不影响任何内容
        if ref_density or ref_motif:
            outline["style_hint"] = {"density": ref_density or "balanced", "motif": ref_motif or "mixed"}
        is_geo = theme in deck_gen.GEO_THEMES
        # 用户没上传照片、又勾了「AI 配图」：几何风 → AI 按 photo_prompts 生成写实照片嵌进图框
        if not photos and ai_bg and is_geo:
            prompts = [p for p in (outline.get("photo_prompts") or []) if isinstance(p, str) and p.strip()][:6]
            if prompts:
                photos = await _gen_deck_spot_photos(prompts, db, user_id)
        outline = _attach_deck_photos(outline, photos)
        bg = None
        if ai_bg and not is_geo:
            detail = bg_detail if bg_detail in ("section", "slide") else "shared"
            raw_kit = await _gen_deck_cover_kit(outline, want_content_bg=(detail == "shared"))
            bg = _persist_bg_kit(db, user_id, raw_kit)
            if detail != "shared":
                # 正文底图改按章节/按页各配一张，挂到每个 slide 的 "bg"；用第一张生成成功的当兜底
                fallback = await _gen_per_page_content_bg(outline, detail, db, user_id)
                if fallback:
                    bg["content"] = fallback
        slides = deck_gen.build_deck(outline, theme, bg)
        _HANDOUT_JOBS[job_id] = {
            "status": "done",
            "result": {"title": outline.get("title") or topic or "演示文稿", "theme": theme, "slides": slides, "outline": outline, "bg": bg},
            "user_id": user_id,
        }
    except HTTPException as e:
        _refund_safely(user_id, ticket)
        _HANDOUT_JOBS[job_id] = {"status": "error", "detail": str(e.detail), "user_id": user_id}
    except Exception as e:  # noqa: BLE001
        _refund_safely(user_id, ticket)
        _HANDOUT_JOBS[job_id] = {"status": "error", "detail": f"生成失败：{e}", "user_id": user_id}
    finally:
        db.close()


def _refund_safely(user_id, ticket):
    from app.database import SessionLocal

    d2 = SessionLocal()
    try:
        u = d2.query(models.User).filter(models.User.id == user_id).first()
        if u:
            billing.refund_ticket(d2, u, ticket)
    except Exception:
        pass
    finally:
        d2.close()


@router.post("/design/deck")
async def design_deck(
    payload: DeckRequest,
    user: models.User = Depends(auth.get_current_user),
    db: Session = Depends(get_db),
):
    """AI 生成 PPT（一期）：主题 → 大纲(+配色) → 排成一套幻灯片。一律走异步 job（轮询
    /design/handout/job/{id}）。ai_bg=True：非几何风生成整页 AI 底图；几何风按主题生成
    4~6 张写实照片嵌进几何图框。计费「AIPPT」。"""
    _require_openlux()
    topic = payload.topic.strip()
    if not topic:
        raise HTTPException(status_code=400, detail="请填写 PPT 主题")
    extra = payload.extra.strip()[:300]
    await _moderate_text(topic + " " + extra)
    n = min(6, max(2, payload.sections))
    photos = _clean_deck_photos(payload.photos)
    ref_pal = [c for c in (payload.palette or []) if re.match(r"^#[0-9a-fA-F]{6}$", str(c))][:5]
    ref_pal = ref_pal if len(ref_pal) == 5 else None
    ticket = billing.consume(db, user, "AIPPT")
    # ai_bg：非几何风 = 生成整页 AI 底图；几何风 = 按主题生成写实照片嵌进几何图框
    ai_bg = bool(payload.ai_bg)
    ref_layouts, ref_density, ref_motif = _clean_ref_hints(
        payload.ref_layouts, payload.ref_density, payload.ref_motif
    )
    bg_detail = payload.bg_detail if payload.bg_detail in ("section", "slide") else "shared"

    # 一律走异步 job：大纲(+可选生图)可能要 1~3 分钟（按页配图能到 10 分钟+），同步返回会被 nginx 网关超时掐断
    job_id = uuid.uuid4().hex
    _HANDOUT_JOBS[job_id] = {"status": "pending", "user_id": user.id}
    _prune_handout_jobs()
    asyncio.create_task(
        _run_deck_job(
            job_id, user.id, ticket, topic, n, payload.theme, extra, ai_bg, "", photos, ref_pal,
            ref_layouts, ref_density, ref_motif, bg_detail,
        )
    )
    return {"jobId": job_id}


_ALLOWED_REF_LAYOUTS = {
    "cards", "list", "timeline", "spoke", "hive", "cycle",
    "matrix", "swot", "gallery", "stats", "bar", "big_number", "quote",
    "tree", "diamond", "bulb", "line", "table",
}


def _clean_ref_hints(layouts, density, motif) -> tuple[list[str], str, str]:
    """前端传来的参考图版式偏好——收窄到白名单，防注入。"""
    ls = [str(x).strip().lower() for x in (layouts or []) if str(x).strip().lower() in _ALLOWED_REF_LAYOUTS][:6]
    d = str(density or "").strip().lower()
    d = d if d in ("airy", "balanced", "packed") else ""
    m = str(motif or "").strip().lower()
    m = m if m in ("hexagon", "circle", "arrow", "wedge", "line", "mixed") else ""
    return ls, d, m


def _shrink_jpeg(image_bytes: bytes, max_px: int = 1024, quality: int = 80) -> tuple[bytes, str]:
    """任意图片 → 缩到 max_px 内的 JPEG 字节。视觉模型调用前统一走这个，
    原图直传 base64 塞 JSON body 弱网易超时。解析失败原样返回。"""
    try:
        im = PILImage.open(io.BytesIO(image_bytes)).convert("RGB")
        if max(im.size) > max_px:
            r = max_px / max(im.size)
            im = im.resize((max(1, round(im.width * r)), max(1, round(im.height * r))), PILImage.LANCZOS)
        b = io.BytesIO()
        im.save(b, "JPEG", quality=quality)
        return b.getvalue(), "image/jpeg"
    except Exception:
        return image_bytes, "image/png"


def _persist_photo(db: Session, user_id: str, image_bytes: bytes) -> str:
    """用户为 PPT 上传的照片：缩到 1600px 转 JPEG 落盘，返回可访问 URL。"""
    data, _ = _shrink_jpeg(image_bytes, max_px=1600, quality=85)
    file_name = f"{uuid.uuid4().hex}.jpg"
    (GENERATED_ASSETS_DIR / file_name).write_bytes(data)
    crud.create_generated_asset(db, user_id, "deck-photo", file_name)
    return f"/api/ai/generated/{file_name}"


async def _transcribe_image_text(image_bytes: bytes, media_type: str) -> str:
    """图片（拍照/截图的备课资料）→ 逐字转录的纯文本，交给大纲模型重组。"""
    data, ct = _shrink_jpeg(image_bytes, max_px=1600, quality=82)
    b64 = base64.b64encode(data).decode()
    res = await _post_openlux(
        f"{OPENLUX_BASE_URL}/chat/completions",
        timeout=120,
        headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
        json={
            "model": "gemini-3-flash-preview",
            "messages": [{
                "role": "user",
                "content": [
                    {"type": "text", "text": (
                        "把这张图片里的所有文字完整、逐字转录出来，保留原有的分段和条目结构。"
                        "不要翻译、不要总结、不要补充说明，只输出文字本身。"
                    )},
                    {"type": "image_url", "image_url": {"url": f"data:{ct};base64,{b64}"}},
                ],
            }],
        },
    )
    if res.status_code >= 400:
        raise HTTPException(status_code=502, detail=f"图片识别失败：{res.status_code} {res.text[:160]}")
    text = (res.json().get("choices", [{}])[0].get("message", {}).get("content") or "").strip()
    if len(text) < 20:
        raise HTTPException(status_code=422, detail="没能从图片里读到足够的文字，换张更清晰的试试")
    return text[:12000]


async def _tag_deck_photos(image_list: list[bytes]) -> list[str]:
    """一批用户上传的照片 → 每张一句中文描述（给幻灯片自动配图用）。一次视觉调用批量处理。"""
    content: list[dict] = [{
        "type": "text",
        "text": (
            "下面是用户为一份 PPT 上传的照片。为每一张写一句 8~24 字的中文描述，"
            "说明画面主体和场景。严格按输入顺序返回一个 JSON 字符串数组，"
            "长度和图片数量一致，只返回数组本身，不要多余说明。"
        ),
    }]
    for raw in image_list:
        data, ct = _shrink_jpeg(raw, max_px=768, quality=76)
        content.append({
            "type": "image_url",
            "image_url": {"url": f"data:{ct};base64,{base64.b64encode(data).decode()}"},
        })
    tags: list[str] = []
    try:
        res = await _post_openlux(
            f"{OPENLUX_BASE_URL}/chat/completions",
            timeout=120,
            headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
            json={"model": "gemini-3-flash-preview", "messages": [{"role": "user", "content": content}]},
        )
        if res.status_code < 400:
            txt = res.json()["choices"][0]["message"]["content"]
            m = re.search(r"\[[\s\S]*\]", txt or "")
            if m:
                tags = [str(x).strip() for x in json.loads(m.group(0)) if str(x).strip()]
    except Exception:
        tags = []
    tags = tags[: len(image_list)]
    while len(tags) < len(image_list):
        tags.append("用户上传的照片")
    return tags


def _clean_deck_photos(photos) -> list[dict]:
    """客户端带回来的 [{url, tag}]——只认我们自己生成资产目录里的文件，防止塞外链。"""
    out: list[dict] = []
    for p in (photos or [])[:12]:
        url = str((p or {}).get("url") or "")
        name = url.rsplit("/", 1)[-1]
        if url.startswith("/api/ai/generated/") and name and "/" not in name and ".." not in name \
                and (GENERATED_ASSETS_DIR / name).is_file():
            out.append({"url": url, "tag": str((p or {}).get("tag") or "用户上传的照片")[:60]})
    return out


def _attach_deck_photos(outline: dict, photos: list[dict]) -> dict:
    """LLM 在 slide 上写的 "image":序号 → 换成真实照片 URL；越界/重复/无效的丢弃。
    顶层 "cover_image":序号 同样处理。photos = [{"url":..,"tag":..}]。"""
    if not photos:
        for sec in outline.get("sections") or []:
            for sl in sec.get("slides") or []:
                sl.pop("image", None)
        outline.pop("cover_image", None)
        return outline
    used: set[int] = set()
    for sec in outline.get("sections") or []:
        for sl in sec.get("slides") or []:
            idx = sl.get("image")
            if isinstance(idx, bool):
                idx = None
            if isinstance(idx, int) and 0 <= idx < len(photos) and idx not in used:
                sl["image"] = photos[idx]["url"]
                used.add(idx)
            else:
                sl.pop("image", None)
            imgs = sl.get("images")
            if isinstance(imgs, list):
                urls = []
                for j in imgs:
                    if isinstance(j, int) and not isinstance(j, bool) and 0 <= j < len(photos) and j not in used:
                        urls.append(photos[j]["url"])
                        used.add(j)
                if len(urls) >= 2:
                    sl["images"] = urls[:3]
                else:
                    sl.pop("images", None)
            else:
                sl.pop("images", None)
    ci = outline.get("cover_image")
    if isinstance(ci, int) and not isinstance(ci, bool) and 0 <= ci < len(photos):
        outline["cover_image"] = photos[ci]["url"]
    else:
        outline.pop("cover_image", None)
    return outline


@router.post("/design/deck/photos")
async def design_deck_photos(
    files: list[UploadFile] = File(...),
    user: models.User = Depends(auth.get_current_user),
    db: Session = Depends(get_db),
):
    """AI PPT 配图：上传若干张真实照片 → 存下来 + 视觉模型逐张打标签，
    返回 [{url, tag}]。前端把它带进 /design/deck 或 /design/deck/material，
    由大纲模型决定哪一页用哪张。不单独计费（生成时按 AIPPT 扣）。"""
    _require_openlux()
    files = files[:12]
    raws: list[bytes] = []
    for f in files:
        raw = await f.read()
        if not raw:
            continue
        if len(raw) > 15 * 1024 * 1024:
            raise HTTPException(status_code=413, detail="单张图片上限 15MB")
        await _check_not_sensitive_document(raw, f.content_type or "image/jpeg", "PPT 配图")
        raws.append(raw)
    if not raws:
        raise HTTPException(status_code=400, detail="没有可用的图片")
    tags = await _tag_deck_photos(raws)
    photos = [{"url": _persist_photo(db, user.id, raw), "tag": tag} for raw, tag in zip(raws, tags)]
    return {"photos": photos}


_DECK_REF_PROMPT = """你在分析一张 PPT 模板/参考图，目的是把它归到我们系统已有的几个通用类别里，用来配置我们自己的模板生成器。我们只学"用了哪几类通用图解、整体多密、主色调"这种最上层的信息，绝不复刻参考图的任何一页版面、坐标、形状或元素数量——那些全部由我们自己的排版引擎独立决定。

只返回一个 JSON 对象，不要多余说明：
{"style":"geo 或 photo 或 plain","palette":["#主色","#强调色","#主色深","#背景浅色","#正文深灰"],"mood":"一句话气质描述","layouts":["从下面固定列表里挑 3~6 个"],"density":"airy 或 balanced 或 packed","motif":"hexagon 或 circle 或 arrow 或 wedge 或 line 或 mixed"}

style：
- "geo"：白底 / 浅底，靠色块、圆弧、环形图、线条等几何图形做装饰
- "photo"：用了实景照片或整幅设计大图当背景 / 主视觉
- "plain"：极简白底，装饰很少

palette：5 个十六进制色，代表整体配色气质（不是逐像素取色）。
mood：例「沉稳的商务深蓝」「科技感的青色调」。

layouts：这套模板"经常出现"的通用图解类型，只能从这个固定列表里选（这些是行业通用的 SmartArt 类别，不是描述某一页）：
  cards（并列要点块）/ list（编号清单）/ timeline（流程时间轴）/ spoke（中心辐射）/ hive（蜂窝六边形群）/ cycle（循环箭头）/ matrix（四象限）/ swot / gallery（多图并排）/ stats（关键指标）/ bar（条形对比）/ big_number（单个大数字）/ quote（金句）/ tree（树状分支图）/ diamond（菱形宫格图标）/ bulb（灯泡放射要点）/ line（折线趋势图）/ table（数据表格）
density：整份看下来页面平均有多满——airy 留白多 / balanced 适中 / packed 信息量大铺得满。
motif：占主导的装饰形状家族——hexagon 六边形 / circle 圆与圆环 / arrow 箭头 / wedge 斜切色块 / line 细线 / mixed 混合。

严禁：描述任何可读文字、精确坐标、精确外形、元素数量、某一页的具体布局。layouts 只是勾选通用类别，不是描述参考图。"""


@router.post("/design/deck/reference")
async def design_deck_reference(
    image: UploadFile = File(...),
    user: models.User = Depends(auth.get_current_user),
    _db: Session = Depends(get_db),
):
    """上传一张喜欢的 PPT 模板/参考图 → 判断它属于我们哪种风格 + 提取配色气质。
    只做「风格归类 + 配色」,不复刻版面。异步 job（轮询 /design/handout/job/{id}），
    结果 {theme, palette, mood, ai_bg}。"""
    _require_openlux()
    raw = await image.read()
    if not raw:
        raise HTTPException(status_code=400, detail="没读到图片")
    ct = image.content_type or "image/png"
    job_id = uuid.uuid4().hex
    _HANDOUT_JOBS[job_id] = {"status": "pending", "user_id": user.id}
    _prune_handout_jobs()
    asyncio.create_task(_run_deck_ref_job(job_id, user.id, raw, ct))
    return {"jobId": job_id}


async def _run_deck_ref_job(job_id, user_id, raw: bytes, ct: str):
    """敏感检查 + 风格分析并行跑，都是视觉调用，串行要 15~40s 常被网关掐断。"""
    try:
        data, ict = _shrink_jpeg(raw, max_px=1024, quality=82)
        b64 = base64.b64encode(data).decode()
        analyze = _post_openlux(
            f"{OPENLUX_BASE_URL}/chat/completions",
            timeout=120,
            headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
            json={
                "model": "gemini-3-flash-preview",
                "messages": [{
                    "role": "user",
                    "content": [
                        {"type": "text", "text": _DECK_REF_PROMPT},
                        {"type": "image_url", "image_url": {"url": f"data:{ict};base64,{b64}"}},
                    ],
                }],
            },
        )
        check_res, res = await asyncio.gather(
            _check_not_sensitive_document(raw, ct, "参考风格"), analyze, return_exceptions=True
        )
        if isinstance(check_res, BaseException):
            raise check_res
        if isinstance(res, BaseException):
            raise res
        if res.status_code >= 400:
            raise HTTPException(status_code=502, detail=f"参考图分析失败：{res.status_code} {res.text[:160]}")
        txt = res.json().get("choices", [{}])[0].get("message", {}).get("content", "")
        m = re.search(r"\{[\s\S]*\}", txt)
        d = json.loads(m.group(0) if m else txt)
        style = str(d.get("style") or "plain").strip().lower()
        pal = [str(c).strip() for c in (d.get("palette") or []) if re.match(r"^#?[0-9a-fA-F]{6}$", str(c).strip())]
        pal = [c if c.startswith("#") else "#" + c for c in pal][:5]
        theme = {"geo": "geoblue", "photo": "techblue"}.get(style, "auto")
        layouts = [
            str(x).strip().lower() for x in (d.get("layouts") or []) if str(x).strip().lower() in _ALLOWED_REF_LAYOUTS
        ][:6]
        density = str(d.get("density") or "").strip().lower()
        density = density if density in ("airy", "balanced", "packed") else "balanced"
        motif = str(d.get("motif") or "").strip().lower()
        motif = motif if motif in ("hexagon", "circle", "arrow", "wedge", "line", "mixed") else "mixed"
        _HANDOUT_JOBS[job_id] = {
            "status": "done",
            "user_id": user_id,
            "result": {
                "theme": theme,
                "palette": pal if len(pal) == 5 else [],
                "mood": str(d.get("mood") or "").strip()[:40],
                "ai_bg": style == "photo",
                "layouts": layouts,
                "density": density,
                "motif": motif,
            },
        }
    except HTTPException as e:
        _HANDOUT_JOBS[job_id] = {"status": "error", "detail": str(e.detail), "user_id": user_id}
    except Exception as e:  # noqa: BLE001
        _HANDOUT_JOBS[job_id] = {"status": "error", "detail": f"参考图分析失败：{e}", "user_id": user_id}


@router.post("/design/deck/material")
async def design_deck_material(
    file: UploadFile | None = File(None),
    text: str = Form(""),
    sections: int = Form(4),
    theme: str = Form("auto"),
    extra: str = Form(""),
    ai_bg: bool = Form(False),
    photos_json: str = Form(""),
    palette_json: str = Form(""),
    ref_hints_json: str = Form(""),
    bg_detail: str = Form("shared"),
    user: models.User = Depends(auth.get_current_user),
    db: Session = Depends(get_db),
):
    """「传资料生成 PPT」：上传 Word/PDF/txt/图片，或直接粘贴长文本 → AI 重组成大纲 →
    排成一套幻灯片。内容全部来自用户资料，不脑补。始终走异步 job（轮询 /design/handout/job/{id}）。
    计费「AIPPT」，跟填主题那条链路一样。"""
    from app import material_extract

    _require_openlux()
    pasted = (text or "").strip()
    material = ""
    if file is not None and (file.filename or ""):
        raw = await file.read()
        if len(raw) > 20 * 1024 * 1024:
            raise HTTPException(status_code=413, detail="文件太大（上限 20MB）")
        ct = (file.content_type or "").lower()
        if ct.startswith("image/") or (file.filename or "").lower().endswith(
            (".png", ".jpg", ".jpeg", ".webp", ".bmp")
        ):
            await _check_not_sensitive_document(raw, ct or "image/png", "生成 PPT")
            material = await _transcribe_image_text(raw, ct or "image/png")
        else:
            try:
                material = material_extract.extract_material(file.filename or "", ct, raw)
            except ValueError as e:
                raise HTTPException(status_code=400, detail=str(e))
    elif pasted:
        if len(pasted) < 20:
            raise HTTPException(status_code=400, detail="粘贴的内容太短，至少写几句")
        material = pasted[: material_extract.MAX_CHARS]
    else:
        raise HTTPException(status_code=400, detail="请上传资料文件或粘贴文字")

    try:
        photos = _clean_deck_photos(json.loads(photos_json) if photos_json.strip() else [])
    except Exception:
        photos = []
    try:
        _p = json.loads(palette_json) if palette_json.strip() else []
        ref_pal = [c for c in _p if re.match(r"^#[0-9a-fA-F]{6}$", str(c))][:5]
        ref_pal = ref_pal if len(ref_pal) == 5 else None
    except Exception:
        ref_pal = None

    try:
        _rh = json.loads(ref_hints_json) if ref_hints_json.strip() else {}
    except Exception:
        _rh = {}
    ref_layouts, ref_density, ref_motif = _clean_ref_hints(
        _rh.get("layouts"), _rh.get("density"), _rh.get("motif")
    )

    await _moderate_text(material[:2000] + " " + extra.strip()[:200])
    n = min(6, max(2, sections))
    ai_bg2 = bool(ai_bg)  # 几何风时 = AI 按主题配图；非几何风 = AI 整页底图
    bg_detail2 = bg_detail if bg_detail in ("section", "slide") else "shared"
    ticket = billing.consume(db, user, "AIPPT")

    job_id = uuid.uuid4().hex
    _HANDOUT_JOBS[job_id] = {"status": "pending", "user_id": user.id}
    _prune_handout_jobs()
    asyncio.create_task(
        _run_deck_job(
            job_id, user.id, ticket, "", n, theme, extra.strip()[:300], ai_bg2, material, photos, ref_pal,
            ref_layouts, ref_density, ref_motif, bg_detail2,
        )
    )
    return {"jobId": job_id}


def _read_generated_asset(src: str) -> bytes | None:
    """/api/ai/generated/<name> → 本地文件字节（PPTX 嵌背景图用）。"""
    name = (src or "").rsplit("/", 1)[-1]
    if not name or "/" in name or ".." in name:
        return None
    p = GENERATED_ASSETS_DIR / name
    return p.read_bytes() if p.is_file() else None


@router.post("/design/deck/pptx")
async def design_deck_pptx(
    payload: DeckPptxRequest,
    _user: models.User = Depends(auth.get_current_user),
):
    """已生成好的幻灯片数据 → PPTX 文件下载。不重新扣次数（生成时已扣）。"""
    from fastapi.responses import StreamingResponse

    data = deck_gen.deck_to_pptx(
        payload.slides, payload.theme, payload.title or "演示文稿", asset_reader=_read_generated_asset)
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": 'attachment; filename="deck.pptx"'},
    )


async def _gen_handout_sections(cat: dict, topic_desc: str, with_content: bool, custom: str = "") -> list[dict]:
    """按分类预设的板块标题让模型填正文，返回 [{heading, items}]；with_content=False 返回 []。"""
    if not with_content:
        return []
    headings = cat["headings"]
    custom_line = f"用户补充要求：{custom}（在保持板块标题和条数不变的前提下尽量体现）。\n" if custom else ""
    content_prompt = (
        f"你在给中小学生的手抄报写内容，主题是「{topic_desc}」（分类：{cat['label']}）。\n"
        f"{custom_line}"
        f"版面已经固定分成这 {len(headings)} 个板块，标题依次是：{'、'.join(headings)}。\n"
        "只返回一个严格的 JSON 对象，不要 markdown 代码块，不要任何多余说明文字，形如：\n"
        '{"sections": [{"heading": "认识安全隐患", "items": ["...", "..."]}]}\n'
        f"要求：sections 数组按上面给定的标题顺序原样输出 {len(headings)} 个，heading 字段必须跟给定标题一字不差；"
        "每个 heading 下面配 3~5 条 items，每条是一句完整的短句（15~30字），内容要准确、具体、适合中小学生阅读，"
        "不能是空话套话；涉及安全/科普类知识点必须准确可靠，不能编造错误常识。"
    )
    chat_res = await _post_openlux(
        f"{OPENLUX_BASE_URL}/chat/completions",
        timeout=60,
        headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
        json={"model": "gemini-3-flash-preview", "messages": [{"role": "user", "content": content_prompt}]},
    )
    if chat_res.status_code >= 400:
        raise HTTPException(status_code=502, detail=f"内容生成失败：{chat_res.status_code} {chat_res.text}")
    raw_content = chat_res.json().get("choices", [{}])[0].get("message", {}).get("content", "")
    match = re.search(r"\{[\s\S]*\}", raw_content)
    try:
        parsed = json.loads(match.group(0) if match else raw_content)
        sections = parsed["sections"]
        if not isinstance(sections, list) or not sections:
            raise ValueError("empty sections")
    except Exception:
        raise HTTPException(status_code=502, detail="手抄报内容生成解析失败，请重试")

    clean_sections: list[dict] = []
    for i, heading in enumerate(headings):
        src = sections[i] if i < len(sections) and isinstance(sections[i], dict) else {}
        items = [str(x).strip() for x in (src.get("items") or []) if str(x).strip()][:6]
        if items:
            clean_sections.append({"heading": heading, "items": items})
    if not clean_sections:
        raise HTTPException(status_code=502, detail="手抄报内容生成为空，请重试")
    return clean_sections


async def _do_handout(db, user, cat: dict, style: dict, border: dict, topic: str, canvas_width: int, canvas_height: int, with_content: bool = True, custom: str = ""):
    topic_desc = topic or cat["label"]

    clean_sections = await _gen_handout_sections(cat, topic_desc, with_content, custom)

    # ── 2. AI 主体插画：集中在左半边，右半 + 顶部留白给文字层叠加；四周可加一圈花边 ──
    # 花边风格：theme 档跟着分类的 motifs 走，none 档不画，其余用预设描述
    if border.get("theme"):
        border_desc = f"{cat['motifs']}"
    else:
        border_desc = border.get("prompt", "")
    if border_desc:
        border_clause = (
            f"画面最外圈画一圈{border_desc}组成的装饰花边，花边宽度约占画布边缘 6%~8%，四个角可以更热闹一些；"
            "花边以内、尤其右侧和顶部，保持纯白留白，不要让花边侵入中间区域。"
        )
    else:
        border_clause = "不要加相框或整圈边框。"

    bg_prompt = (
        f"儿童手抄报的主体插画，主题「{topic_desc}」。{style['prompt']}。{cat['palette']}。\n"
        f"画面主体：{cat['scene']}。"
        f"周围点缀：{cat['motifs']}，元素要丰富、错落有致。\n"
        "构图要求：主体插画和人物全部集中在画面左侧约 45% 的范围内，画面右侧 55% 和顶部 25% 必须是纯白或极淡的底色，"
        "不能有任何人物、图案、文字——那些区域后面要叠加排版好的文字。"
        f"{border_clause}"
        "白色背景，不要画标题文字。"
        + (f"\n用户额外要求（在不破坏上面构图/留白规则的前提下尽量满足）：{custom}" if custom else "")
    )
    colored_src = None
    lineart_src = None
    asset_id = None
    for _attempt in range(2):
        try:
            gen_res = await _post_openlux(
                f"{OPENLUX_BASE_URL}/images/generations",
                timeout=150,
                headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
                json={"model": "gpt-image-2", "prompt": bg_prompt, "n": 1, "size": _nearest_image_size(canvas_width, canvas_height)},
            )
            if gen_res.status_code >= 400:
                continue
            try:
                image_bytes = await _extract_openai_image_bytes(gen_res.json(), "手抄报插画")
            except HTTPException:
                continue
            colored_asset = _persist_asset_bytes(db, user.id, "handout-colored", image_bytes)
            asset_id = colored_asset.id
            colored_src = f"/api/ai/generated/{colored_asset.file_name}"
            # 线稿版：OpenCV 从彩色版提，构图跟彩色版完全一致——家长照着彩色的给孩子涂
            try:
                lineart_asset = _persist_asset_bytes(db, user.id, "handout-lineart", _to_lineart(image_bytes))
                lineart_src = f"/api/ai/generated/{lineart_asset.file_name}"
            except Exception:
                pass
            break
        except Exception:
            continue

    # ── 3. 排版：顶部艺术大标题 + 右侧文字板块（后端确定性算好，前端叠图上去）────
    layout = layout_presets.build_handout(
        canvas_width, canvas_height, topic_desc, clean_sections, cat["colors"], with_content=with_content,
    )

    return {
        "coloredSrc": colored_src,
        "lineartSrc": lineart_src,
        "backgroundSrc": colored_src,  # 兼容旧前端字段
        "background": layout["background"],
        "elements": layout["elements"],
        "title": topic_desc,
        "sections": clean_sections,
        "colors": cat["colors"],
        "assetId": asset_id,
    }


_LIST_ELEMENTS_PROMPT = """这是一张儿童手抄报。列出画面里所有独立的图画元素——人物、动物、
植物、道具、装饰小图标等，不要列文字、不要列四周花边、不要列整张纸。
只返回一个 JSON 数组，不要 markdown、不要多余说明，每项形如：
{"name": "敬礼的少年", "box": [ymin, xmin, ymax, xmax]}
name 是 2~8 个字的简短中文；box 坐标 0~1000 整数，(0,0) 在左上角。
按画面里的重要性/大小排序，最多 9 个，别把两个物体合成一项。"""


async def _list_handout_elements(image_b64: str, max_n: int = 9) -> list[dict]:
    """视觉模型列出整图里的元素名 + 大致位置。失败返回 []。"""
    try:
        res = await _post_openlux(
            f"{OPENLUX_BASE_URL}/chat/completions",
            timeout=90,
            headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
            json={
                "model": "gemini-3-flash-preview",
                "messages": [{"role": "user", "content": [
                    {"type": "text", "text": _LIST_ELEMENTS_PROMPT},
                    {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{image_b64}"}},
                ]}],
            },
        )
    except Exception:
        return []
    if res.status_code >= 400:
        return []
    raw = res.json().get("choices", [{}])[0].get("message", {}).get("content", "")
    m = re.search(r"\[[\s\S]*\]", raw)
    if not m:
        return []
    try:
        items = json.loads(m.group(0))
    except Exception:
        return []
    out: list[dict] = []
    for it in items:
        if not isinstance(it, dict):
            continue
        name = str(it.get("name") or it.get("label") or "").strip()[:16]
        box = it.get("box") or it.get("box_2d") or it.get("bbox") or [200, 200, 500, 500]
        try:
            ymin, xmin, ymax, xmax = [max(0.0, min(1000.0, float(v))) for v in box]
        except (TypeError, ValueError):
            ymin, xmin, ymax, xmax = 200, 200, 500, 500
        if not name:
            continue
        out.append({"name": name, "box": [xmin, ymin, xmax, ymax]})
        if len(out) >= max_n:
            break
    return out


async def _decompose_to_layers(db, user, full_bytes: bytes, W: int, H: int, style_hint: str, *, make_bg: bool = True):
    """把一张整图拆成图层：视觉模型列元素 → gemini 图像模型「照画风单独重画」每个成透明贴纸
    （不是裁切，所以每个都干净）→ 可选再出一张"只有花边、中间留白"的底图。
    返回 {elImages, bgSrc, fullAssetId}。"""
    full_b64 = base64.b64encode(full_bytes).decode()
    full_asset = _persist_asset_bytes(db, user.id, "handout-full", full_bytes)
    specs = (await _list_handout_elements(full_b64))[:8]
    sem = asyncio.Semaphore(4)  # openlux 上游对 gemini image 有并发限制

    async def _make_sticker(spec: dict):
        prompt = (
            f"参考这张图的画风和配色，单独画一个「{spec['name']}」，{style_hint}，粗黑描边，"
            "只有这一个物体、完整居中、占满画面，纯透明背景，不要文字、不要边框、不要其他物体、不要地面阴影。"
        )
        async with sem:
            try:
                b = await _gemini_image(prompt, ref_b64=full_b64, attempts=3, timeout=140)
            except Exception:
                return None
        try:
            a = _persist_asset_bytes(db, user.id, "handout-element", b)
        except Exception:
            return None
        x0, y0, x1, y1 = spec["box"]
        return {
            "type": "image",
            "x": round(x0 / 1000 * W),
            "y": round(y0 / 1000 * H),
            "width": max(40, round((x1 - x0) / 1000 * W)),
            "height": max(40, round((y1 - y0) / 1000 * H)),
            "src": f"/api/ai/generated/{a.file_name}",
        }

    async def _make_bg():
        if not make_bg:
            return None
        prompt = (
            "参考这张图的画风，只画四周一圈的装饰花边（跟原图同样的元素和配色），"
            "画面正中间大片完全留白，不要画任何人物/动物/道具/主体插画。纯白纸底，输出 PNG。"
        )
        async with sem:
            try:
                return await _gemini_image(prompt, ref_b64=full_b64, attempts=2, timeout=140)
            except Exception:
                return None

    sticker_results, bg_bytes = await asyncio.gather(
        asyncio.gather(*[_make_sticker(s) for s in specs]),
        _make_bg(),
    )
    el_images = [r for r in sticker_results if r]
    bg_src = None
    if bg_bytes:
        bg_asset = _persist_asset_bytes(db, user.id, "handout-bg", bg_bytes)
        bg_src = f"/api/ai/generated/{bg_asset.file_name}"
    return {"elImages": el_images, "bgSrc": bg_src, "fullAssetId": full_asset.id,
            "fullName": full_asset.file_name}


async def _do_handout_layered(db, user, cat: dict, style: dict, border: dict, topic: str, W: int, H: int, with_content: bool = True, custom: str = ""):
    """AI 出整张手抄报当参考 → 拆成一堆可拖动/可提示词替换的透明贴纸图层 + 标题/文字层。"""
    topic_desc = topic or cat["label"]
    clean_sections = await _gen_handout_sections(cat, topic_desc, with_content, custom)

    border_desc = cat["motifs"] if border.get("theme") else border.get("prompt", "")
    border_clause = f"画面四周画一圈{border_desc}组成的花边（约占边缘 6%）。" if border_desc else ""
    full_prompt = (
        f"儿童手抄报整张画面，主题「{topic_desc}」。{style['prompt']}。{cat['palette']}。\n"
        f"画面里有：{cat['scene']}；另外散布这些装饰元素：{cat['motifs']}。\n"
        f"{border_clause}"
        "主体元素集中在画面左侧 60% 和中间，右侧 40% 保持纯白留白。不要画任何文字、横线或方格。"
        + (f"\n用户额外要求（不要画文字，其余尽量满足）：{custom}" if custom else "")
    )
    raw = await _gen_image_bytes(full_prompt, _nearest_image_size(W, H), attempts=2, timeout=170)

    dec = await _decompose_to_layers(db, user, raw, W, H, f"{style['prompt']}，{cat['palette']}", make_bg=True)

    bg_color = "#ffffff" if dec["bgSrc"] else layout_presets._tint(cat["colors"][0], 0.92)
    layout = layout_presets.build_handout(W, H, topic_desc, clean_sections, cat["colors"], with_content=with_content)
    elements: list[dict] = []
    if dec["bgSrc"]:
        elements.append({"type": "image", "x": 0, "y": 0, "width": W, "height": H, "src": dec["bgSrc"]})
    elements += dec["elImages"]
    elements += layout["elements"]

    return {
        "background": bg_color,
        "elements": elements,
        "title": topic_desc,
        "sections": clean_sections,
        "colors": cat["colors"],
        "fullSrc": f"/api/ai/generated/{dec['fullName']}",
        "elementCount": len(dec["elImages"]),
        "assetId": dec["fullAssetId"],
    }


_ANALYZE_HANDOUT_PROMPT = """这是一张手抄报/黑板报图片。请拆解它的结构，只返回一个 JSON 对象，
不要 markdown、不要多余说明，形如：
{
  "illustrations": [{"name": "太阳", "box": [ymin,xmin,ymax,xmax]}],
  "textBoxes": [{"heading": "地球需要我们", "body": "地球是我们唯一的家园，空气、水...", "box": [ymin,xmin,ymax,xmax], "color": "#5aa832"}]
}
- illustrations：画面里的图画元素（人物/动物/植物/道具/云朵/装饰图标等），不含文字、不含四周花边、不含整块背景。最多 12 个。
- textBoxes：带颜色边框的文字方框，heading 是小标题、body 是框里的正文（原样抄写，去掉换行），color 是这个框边框的主色（hex）。没有文字方框就返回空数组。
- 所有 box 坐标是 0~1000 的整数，(0,0) 在左上角。"""


async def _analyze_handout_image(image_b64: str) -> dict:
    try:
        res = await _post_openlux(
            f"{OPENLUX_BASE_URL}/chat/completions",
            timeout=90,
            headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
            json={
                "model": "gemini-3-flash-preview",
                "messages": [{"role": "user", "content": [
                    {"type": "text", "text": _ANALYZE_HANDOUT_PROMPT},
                    {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{image_b64}"}},
                ]}],
            },
        )
        if res.status_code >= 400:
            return {"illustrations": [], "textBoxes": []}
        raw = res.json().get("choices", [{}])[0].get("message", {}).get("content", "")
        m = re.search(r"\{[\s\S]*\}", raw)
        data = json.loads(m.group(0)) if m else {}
    except Exception:
        return {"illustrations": [], "textBoxes": []}
    return {
        "illustrations": data.get("illustrations") or [],
        "textBoxes": data.get("textBoxes") or [],
    }


def _box_to_px(box, iw: int, ih: int, pad: int = 0):
    """gemini 的 [ymin,xmin,ymax,xmax] 0~1000 → 像素 (x0,y0,x1,y1)，带 pad、裁到画布内。"""
    try:
        ymin, xmin, ymax, xmax = [max(0.0, min(1000.0, float(v))) for v in box]
    except (TypeError, ValueError):
        return None
    x0 = max(0, int(xmin / 1000 * iw) - pad)
    y0 = max(0, int(ymin / 1000 * ih) - pad)
    x1 = min(iw, int(xmax / 1000 * iw) + pad)
    y1 = min(ih, int(ymax / 1000 * ih) + pad)
    return (x0, y0, x1, y1) if x1 - x0 > 8 and y1 - y0 > 8 else None


async def _do_decompose(db, user, image_bytes: bytes, W: int, H: int, style_key: str):
    """把用户上传的现成手抄报图拆成可编辑图层：
    - 图画元素：直接从原图裁下来 + 抠白底（忠于原图、不重画、不变形），并在底图上抹白
    - 文字方框：换成"彩色圆角边框 + 可编辑小标题 + 可编辑正文（预填原文）"，边框颜色可改
    - 底图：原图去掉被裁走的图画元素，花边/标题保留
    """
    src_img = cv2.imdecode(np.frombuffer(image_bytes, np.uint8), cv2.IMREAD_COLOR)
    if src_img is None:
        raise HTTPException(status_code=400, detail="图片打不开，换一张")
    ih, iw = src_img.shape[:2]
    b64 = base64.b64encode(image_bytes).decode()
    analysis = await _analyze_handout_image(b64)

    sx, sy = W / iw, H / ih
    bg_img = src_img.copy()
    el_images: list[dict] = []

    for ill in analysis["illustrations"][:12]:
        px = _box_to_px(ill.get("box") or [], iw, ih, pad=round(min(iw, ih) * 0.008))
        if not px:
            continue
        x0, y0, x1, y1 = px
        if (x1 - x0) * (y1 - y0) > 0.4 * iw * ih:  # 太大八成是框错了整块
            continue
        ok, cbuf = cv2.imencode(".png", src_img[y0:y1, x0:x1])
        if not ok:
            continue
        try:
            cut = _cutout_white_bg(cbuf.tobytes())
        except Exception:
            cut = cbuf.tobytes()
        a = _persist_asset_bytes(db, user.id, "handout-element", cut)
        el_images.append({
            "type": "image",
            "x": round(x0 * sx), "y": round(y0 * sy),
            "width": max(30, round((x1 - x0) * sx)), "height": max(30, round((y1 - y0) * sy)),
            "src": f"/api/ai/generated/{a.file_name}",
        })
        cv2.rectangle(bg_img, (x0, y0), (x1, y1), (255, 255, 255), -1)

    ok, bgbuf = cv2.imencode(".png", bg_img)
    bg_asset = _persist_asset_bytes(db, user.id, "handout-bg", bgbuf.tobytes() if ok else image_bytes)
    full_asset = _persist_asset_bytes(db, user.id, "handout-full", image_bytes)

    text_els: list[dict] = []
    for tb in analysis["textBoxes"][:8]:
        px = _box_to_px(tb.get("box") or [], iw, ih, pad=round(min(iw, ih) * 0.012))
        if not px:
            continue
        x0, y0, x1, y1 = [v for v in px]
        cx0, cy0 = round(x0 * sx), round(y0 * sy)
        cw, ch = round((x1 - x0) * sx), round((y1 - y0) * sy)
        color = str(tb.get("color") or "#2563eb").strip()
        if not re.match(r"^#[0-9a-fA-F]{6}$", color):
            color = "#2563eb"
        heading = str(tb.get("heading") or "").strip()[:16]
        body = str(tb.get("body") or "").strip()
        scale = max(0.7, min(2.2, cw / 320))
        fs_head = round(20 * scale)
        fs_body = round(15 * scale)
        pad = round(14 * scale)
        # 不透明淡色底 + 同色描边，盖住原图那个框
        text_els.append({"type": "rect", "x": cx0, "y": cy0, "width": cw, "height": ch,
                         "fill": layout_presets._tint(color, 0.9), "rx": round(16 * scale),
                         "stroke": color, "strokeWidth": max(2, round(2.4 * scale))})
        if heading:
            hw = min(cw - 2 * pad, round(len(heading) * fs_head * 1.2 + fs_head * 1.8))
            hh = round(fs_head * 1.7)
            text_els.append({"type": "rect", "x": cx0 + pad, "y": cy0 + pad, "width": hw, "height": hh,
                             "fill": color, "rx": round(hh / 2)})
            text_els.append({"type": "text", "x": cx0 + pad, "y": cy0 + pad + round((hh - fs_head) / 2 - 1),
                             "width": hw, "text": heading, "fontSize": fs_head, "color": "#ffffff",
                             "fontWeight": "bold", "align": "center"})
        text_els.append({"type": "text", "x": cx0 + pad, "y": cy0 + pad + (round(fs_head * 1.7) + round(10 * scale) if heading else 0),
                         "width": cw - 2 * pad, "text": body or "（点这里改文字）", "fontSize": fs_body,
                         "color": "#374151"})

    elements: list[dict] = [{"type": "image", "x": 0, "y": 0, "width": W, "height": H,
                             "src": f"/api/ai/generated/{bg_asset.file_name}"}]
    elements += el_images
    elements += text_els

    return {
        "background": "#ffffff",
        "elements": elements,
        "title": "",
        "sections": [],
        "colors": ["#dc2626", "#2563eb"],
        "fullSrc": f"/api/ai/generated/{full_asset.file_name}",
        "elementCount": len(el_images) + len(analysis["textBoxes"][:8]),
        "assetId": full_asset.id,
    }


async def _run_decompose_job(job_id, user_id, ticket, image_bytes, W, H, style_key):
    from app.database import SessionLocal

    db = SessionLocal()
    try:
        user = db.query(models.User).filter(models.User.id == user_id).first()
        if user is None:
            raise RuntimeError("用户不存在")
        result = await _do_decompose(db, user, image_bytes, W, H, style_key)
        _HANDOUT_JOBS[job_id] = {"status": "done", "result": result, "user_id": user_id}
    except Exception as e:  # noqa: BLE001
        try:
            user = db.query(models.User).filter(models.User.id == user_id).first()
            if user:
                billing.refund_ticket(db, user, ticket)
        except Exception:
            pass
        detail = str(e.detail) if isinstance(e, HTTPException) else f"拆解失败：{e}"
        _HANDOUT_JOBS[job_id] = {"status": "error", "detail": detail, "user_id": user_id}
    finally:
        db.close()


@router.post("/design/decompose")
async def design_decompose(
    image: UploadFile = File(...),
    canvas_width: int = Form(...),
    canvas_height: int = Form(...),
    style: str = Form("color"),
    user: models.User = Depends(auth.get_current_user),
    db: Session = Depends(get_db),
):
    """上传一张现成的整图（豆包/别处生成的手抄报），拆成可单独拖动/替换的图层。异步，轮询同 handout。"""
    _require_openlux()
    image_bytes = await image.read()
    await _check_not_sensitive_document(image_bytes, image.content_type or "image/png", "图片拆解")
    ticket = billing.consume(db, user, "手抄报拆分")
    job_id = uuid.uuid4().hex
    _HANDOUT_JOBS[job_id] = {"status": "pending", "user_id": user.id}
    _prune_handout_jobs()
    asyncio.create_task(
        _run_decompose_job(job_id, user.id, ticket, image_bytes, canvas_width, canvas_height, style)
    )
    return {"jobId": job_id}


MAX_CONVERT_SLIDES = 12


def _extract_full_bleed_slide_images(pptx_bytes: bytes, min_coverage: float = 0.8) -> list[bytes]:
    """"截图型 PPT"（每页就是一张铺满整页的图片，截图/图片拼的、没法编辑那种）→
    每页的图片字节，按页序排好。要求每一页都得是这种整页图——只要有一页本来就是正常的可编辑
    文字页，就报错说第几页不满足，不做"部分转换"（语义含糊，不如让用户明确知道这工具是干嘛的）。"""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    try:
        prs = Presentation(io.BytesIO(pptx_bytes))
    except Exception as e:  # noqa: BLE001
        raise ValueError(f"PPT 文件打不开：{e}") from e

    slides = list(prs.slides)
    if not slides:
        raise ValueError("这份 PPT 没有幻灯片")
    if len(slides) > MAX_CONVERT_SLIDES:
        raise ValueError(f"最多支持 {MAX_CONVERT_SLIDES} 页（这份有 {len(slides)} 页）")

    slide_area = prs.slide_width * prs.slide_height
    out: list[bytes] = []
    for i, slide in enumerate(slides, 1):
        best = None
        best_area = 0
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                area = shape.width * shape.height
                if area > best_area:
                    best_area, best = area, shape
        if best is None or slide_area <= 0 or best_area / slide_area < min_coverage:
            raise ValueError(f"第 {i} 页不是整页截图——这个功能只处理「每页都是一张图」的 PPT")
        out.append(best.image.blob)
    return out


async def _do_pptx_convert(db, user, slide_images: list[bytes], style_key: str) -> dict:
    """逐页复用手抄报的拆图层引擎（_do_decompose），拼回一套「AI PPT」结果结构——
    跟 deck_gen 的 slides 数组（background + elements + w/h）完全同格式，下载 PPTX
    直接走已有的 deck_gen.deck_to_pptx，不用另写导出器。"""
    slides: list[dict] = []
    for img_bytes in slide_images:
        await _check_not_sensitive_document(img_bytes, "image/png", "PPT 转可编辑")
        dec = await _do_decompose(db, user, img_bytes, deck_gen.W, deck_gen.H, style_key)
        slides.append({
            "background": dec["background"],
            "elements": dec["elements"],
            "w": deck_gen.W,
            "h": deck_gen.H,
        })
    return {"title": "转换后的 PPT", "theme": "auto", "slides": slides}


async def _run_pptx_convert_job(job_id, user_id, ticket, slide_images: list[bytes], style_key: str):
    from app.database import SessionLocal

    db = SessionLocal()
    try:
        user = db.query(models.User).filter(models.User.id == user_id).first()
        if user is None:
            raise RuntimeError("用户不存在")
        result = await _do_pptx_convert(db, user, slide_images, style_key)
        _HANDOUT_JOBS[job_id] = {"status": "done", "result": result, "user_id": user_id}
    except Exception as e:  # noqa: BLE001
        try:
            user = db.query(models.User).filter(models.User.id == user_id).first()
            if user:
                billing.refund_ticket(db, user, ticket)
        except Exception:
            pass
        detail = str(e.detail) if isinstance(e, HTTPException) else f"转换失败：{e}"
        _HANDOUT_JOBS[job_id] = {"status": "error", "detail": detail, "user_id": user_id}
    finally:
        db.close()


@router.post("/design/pptx-to-editable")
async def design_pptx_to_editable(
    file: UploadFile = File(...),
    style: str = Form("color"),
    user: models.User = Depends(auth.get_current_user),
    db: Session = Depends(get_db),
):
    """上传"图片型 PPT"——每页都是截图/图片拼的、在 PowerPoint 里选不中文字那种——
    逐页拆图层（复用手抄报拆分引擎）重新拼成一套原生形状/文本框的可编辑 PPTX。
    异步，轮询同 handout（/design/handout/job/{id}）。"""
    _require_openlux()
    raw = await file.read()
    if len(raw) > 40 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="文件太大（上限 40MB）")
    try:
        slide_images = _extract_full_bleed_slide_images(raw)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

    ticket = billing.consume(db, user, "PPT转可编辑")
    job_id = uuid.uuid4().hex
    _HANDOUT_JOBS[job_id] = {"status": "pending", "user_id": user.id}
    _prune_handout_jobs()
    asyncio.create_task(_run_pptx_convert_job(job_id, user.id, ticket, slide_images, style))
    return {"jobId": job_id}


@router.get("/assets")
def list_generated_assets(
    category: str | None = None,
    user: models.User = Depends(auth.get_current_user),
    db: Session = Depends(get_db),
):
    """"素材"面板用——只列当前登录用户自己生成/保存的图，私有，不是公共素材库。"""
    rows = crud.list_generated_assets(db, user.id, category)
    return {
        "list": [
            {"id": r.id, "category": r.category, "url": f"/api/ai/generated/{r.file_name}", "createdAt": r.created_at.isoformat()}
            for r in rows
        ]
    }


@router.delete("/assets/{asset_id}")
def delete_generated_asset(
    asset_id: str,
    user: models.User = Depends(auth.get_current_user),
    db: Session = Depends(get_db),
):
    asset = crud.get_generated_asset(db, asset_id)
    if not asset or asset.user_id != user.id:
        raise HTTPException(status_code=404, detail="素材不存在")
    file_path = GENERATED_ASSETS_DIR / asset.file_name
    if file_path.exists():
        file_path.unlink()
    crud.delete_generated_asset(db, asset_id)
    return {"deleted": True}


# 抠图 / 上色这类要现场加载几百 MB~1GB 模型的本地子进程，全局一次只放一个进去跑——
# 这台机器只有 2GB（实测空载可用 ~1.4GB），两个大模型进程同时加载会把内存打爆。
# 并发请求排队等前一个跑完（每个也就几秒），对当前访问量完全无感。
_HEAVY_MODEL_SEM = asyncio.Semaphore(1)


def _run_bg_removal_subprocess(image_bytes: bytes, edge: str = "soft") -> subprocess.CompletedProcess:
    """用同步 subprocess.run（不是 asyncio.create_subprocess_exec）——后者在 Windows 上
    uvicorn --reload 强制用的 SelectorEventLoop 下会直接抛 NotImplementedError（Windows
    独有的坑，asyncio 文档里写明 selector loop 不支持子进程；Linux 生产环境不受影响，
    但这样写本地 Windows 开发环境也能跑通同一条代码路径，不用靠"生产是 Linux 应该没事"硬赌）。
    外层用 asyncio.to_thread 扔到线程池，不阻塞事件循环。edge 透传给 worker 决定边缘处理档位。"""
    from app.worker_cmd import worker_argv

    return subprocess.run(
        worker_argv("bg_removal", edge),
        input=image_bytes,
        capture_output=True,
        timeout=240,  # 首次要下 ~178MB 模型，慢一点；之后走缓存几秒钟
    )


@router.post("/background-removal")
async def background_removal(
    image: UploadFile = File(...),
    edge: str = Form("soft"),
    _user: models.User = Depends(auth.get_current_user),
):
    """本地跑 rembg（isnet-general-use 模型）本身不依赖任何第三方 key，但合规检查需要调用
    openlux 的视觉模型，所以这个接口现在也要求 OPENLUX_API_KEY 配置好——这是"所有工具统一走
    一层敏感文件检测"这条要求带来的必然依赖，不是抠图本身需要。放到独立子进程里跑（见
    bg_removal_worker.py）——模型稳定态占约 1GB RSS，不能常驻在主 uvicorn 进程里，
    不然这台 2 核 2GB 的机器迟早被这个功能拖垮。首次调用要下载模型（~178MB），
    会明显慢一次，之后走本地缓存就快了。edge=soft 保细节，edge=hard 出照相馆硬边。"""
    _require_openlux()
    image_bytes = await image.read()
    await _check_not_sensitive_document(image_bytes, image.content_type or "image/png", "AI 抠图")
    edge_mode = edge if edge in ("soft", "hard") else "soft"
    async with _HEAVY_MODEL_SEM:
        try:
            proc = await asyncio.to_thread(_run_bg_removal_subprocess, image_bytes, edge_mode)
        except subprocess.TimeoutExpired:
            raise HTTPException(status_code=504, detail="抠图处理超时，请重试")
    if proc.returncode != 0:
        raise HTTPException(status_code=502, detail=f"抠图处理失败：{proc.stderr.decode(errors='ignore')[:500]}")
    b64 = base64.b64encode(proc.stdout).decode()
    return {"data": [{"b64_json": b64}]}


def _run_colorize_subprocess(image_bytes: bytes, saturation: float) -> subprocess.CompletedProcess:
    """同 _run_bg_removal_subprocess：同步 subprocess.run + asyncio.to_thread 包一层，
    绕开 asyncio 子进程在 Windows selector loop 上的坑，本地/生产同一条代码路径。"""
    from app.worker_cmd import worker_argv

    return subprocess.run(
        worker_argv("colorize", str(saturation)),
        input=image_bytes,
        capture_output=True,
        timeout=170,
    )


@router.post("/colorize")
async def colorize_photo(
    image: UploadFile = File(...),
    saturation: float = Form(1.0),
    _user: models.User = Depends(auth.get_current_user),
):
    """老照片上色（黑白 -> 彩色）。跟"黑白遗像"是相反方向的两件事——遗像是彩色转黑白 + 适配
    相框尺寸，这个是黑白转彩色。纯本地推理（DDColor large int8 量化 ONNX，走 onnxruntime
    CPU），不调用 openlux，也不接敏感文件检测——上色是"修复/还原"不是"篡改"，跟扫描件/
    提字一个定性（你没法靠给黑白照片上色去伪造证件，证件本来就是彩色的）。放独立子进程跑，
    推理峰值约 600MB，处理完立刻释放。首次调用现场下载模型（~235MB），会明显慢一次。
    worker 会先做一遍轻度对比修复（翻拍老照片普遍偏灰，不修的话上色效果很淡）。
    saturation：>1 提饱和（模型输出偏保守时用），范围 0.5~2.0。"""
    image_bytes = await image.read()
    sat = min(2.0, max(0.5, saturation))
    async with _HEAVY_MODEL_SEM:
        try:
            proc = await asyncio.to_thread(_run_colorize_subprocess, image_bytes, sat)
        except subprocess.TimeoutExpired:
            raise HTTPException(status_code=504, detail="上色处理超时，请重试")
    if proc.returncode != 0:
        raise HTTPException(status_code=502, detail=f"上色处理失败：{proc.stderr.decode(errors='ignore')[:500]}")
    b64 = base64.b64encode(proc.stdout).decode()
    return {"data": [{"b64_json": b64, "mime": "image/jpeg"}]}


_FACE_CASCADE = cv2.CascadeClassifier(cv2.data.haarcascades + "haarcascade_frontalface_default.xml")


@router.post("/detect-face")
async def detect_face(
    image: UploadFile = File(...),
    _user: models.User = Depends(auth.get_current_user),
):
    """证件照制作专用：本地跑 OpenCV Haar 级联检测人脸位置，纯本地算法不调用 openlux，
    零额外 AI 成本。返回人脸框相对图片宽高的 0~1 比例坐标，跟 reference-to-asset 的框选坐标
    是同一套单位约定，前端好复用。

    minSize 按图片自身尺寸的比例算，不能用固定像素——实测过：固定 100px 在 4032×3024 的真实
    照片上能准确避开衣服纹理这类误判，但换到 128×121 的小图上 100px 比整张图还大，直接测不出
    人脸。多张脸（真实照片背景纹理偶尔会被误判成好几个小"脸"）取面积最大的一个，实测这个策略
    在真实测试照片上（一张 128×121 小头像 + 一张 4032×3024 手机原图，原图上有 7 个误判但都
    明显小于真脸）都精确只留下真脸。检测不到人脸时返回 null，前端退回原来"整图占比"的粗略估算。"""
    image_bytes = await image.read()
    arr = np.frombuffer(image_bytes, dtype=np.uint8)
    img = cv2.imdecode(arr, cv2.IMREAD_COLOR)
    if img is None:
        raise HTTPException(status_code=400, detail="无法解析上传的图片")
    h, w = img.shape[:2]
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    min_size = (max(30, int(w * 0.08)), max(30, int(h * 0.08)))
    faces = _FACE_CASCADE.detectMultiScale(gray, scaleFactor=1.1, minNeighbors=6, minSize=min_size)
    if len(faces) == 0:
        return {"face": None}
    fx, fy, fw, fh = max(faces, key=lambda f: int(f[2]) * int(f[3]))
    return {"face": {"x": float(fx) / w, "y": float(fy) / h, "width": float(fw) / w, "height": float(fh) / h}}


_EYE_CASCADE = cv2.CascadeClassifier(cv2.data.haarcascades + "haarcascade_eye.xml")


@router.post("/fix-red-eye")
async def fix_red_eye(
    image: UploadFile = File(...),
    _user: models.User = Depends(auth.get_current_user),
):
    """去红眼——纯本地 OpenCV，跟人脸检测同一套零成本算法，不调用 openlux。先用人脸检测框
    缩小眼部检测的搜索范围（直接在整张图上跑眼部级联，背景纹理误判率明显更高——这是复用
    人脸检测代码路径的直接原因，不是巧合）；在人脸区域里跑眼部检测，对每个检测到的眼睛，
    把"红色通道明显高于绿蓝通道且不算太暗"的像素（闪光灯反射视网膜的典型特征）单独替换成
    去掉红色分量后的灰阶，只动红色通道，不是整片涂黑——保留瞳孔原有的明暗结构。

    没检测到人脸时退回整张图找眼睛（不直接放弃），检测不到眼睛就原样返回，不报错——
    没有真实红眼样张走过真实场景测试，效果依赖真实红眼照片实测反馈。"""
    image_bytes = await image.read()
    arr = np.frombuffer(image_bytes, dtype=np.uint8)
    img = cv2.imdecode(arr, cv2.IMREAD_COLOR)
    if img is None:
        raise HTTPException(status_code=400, detail="无法解析上传的图片")
    h, w = img.shape[:2]
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    face_min_size = (max(30, int(w * 0.08)), max(30, int(h * 0.08)))
    faces = _FACE_CASCADE.detectMultiScale(gray, scaleFactor=1.1, minNeighbors=6, minSize=face_min_size)
    if len(faces) > 0:
        fx, fy, fw, fh = max(faces, key=lambda f: int(f[2]) * int(f[3]))
    else:
        fx, fy, fw, fh = 0, 0, w, h

    face_gray = gray[fy : fy + fh, fx : fx + fw]
    eye_min_size = (max(15, int(fw * 0.12)), max(15, int(fh * 0.12)))
    eyes = _EYE_CASCADE.detectMultiScale(face_gray, scaleFactor=1.1, minNeighbors=8, minSize=eye_min_size)

    fixed_count = 0
    for ex, ey, ew, eh in eyes:
        ax, ay = fx + int(ex), fy + int(ey)
        region = img[ay : ay + eh, ax : ax + ew].astype(np.int16)
        b, g, r = region[:, :, 0], region[:, :, 1], region[:, :, 2]
        red_mask = (r > g * 1.4) & (r > b * 1.4) & (r > 60)
        # 至少 2% 的区域被判定为偏红才当作真的红眼处理——实测过一只完全正常的眼睛也会有
        # 零星 1~2 个像素误判（反光高光点），只看"有没有"会把正常眼睛也算进"修复了几只"，
        # 卡个比例门槛能把这类噪声排除掉
        if red_mask.mean() < 0.02:
            continue
        avg_gb = ((g + b) / 2).astype(np.uint8)
        red_channel = img[ay : ay + eh, ax : ax + ew, 2]
        red_channel[red_mask] = avg_gb[red_mask]
        fixed_count += 1

    success, buf = cv2.imencode(".png", img)
    if not success:
        raise HTTPException(status_code=500, detail="图片编码失败")
    b64 = base64.b64encode(buf.tobytes()).decode()
    return {"data": [{"b64_json": b64}], "eyesFixed": fixed_count}


@router.post("/convert-cmyk")
async def convert_cmyk(
    image: UploadFile = File(...),
    _user: models.User = Depends(auth.get_current_user),
):
    """把浏览器画出来的 RGB 图转成 CMYK JPEG——只有真的要送商业印刷厂（胶印/丝网印这类需要
    CMYK 四色分色的场景）才用得上。Canvas API 只能画 RGB，这一步必须经后端：Pillow 的
    convert('CMYK') 是直接的数学换算，不带 ICC 色彩管理描述文件，转出来的颜色（尤其肤色）
    会比印刷厂用专业软件转的偏灰/偏暗——只是格式对了，不代表色彩准。普通冲印店/家用打印机
    认 RGB 就够，不需要这一步；生成的 CMYK JPEG 在浏览器里直接用 <img> 预览可能显示异常
    （多数浏览器的 JPEG 解码器按 RGB/YCbCr 假设，不认 CMYK JPEG），只适合直接下载后
    在 Photoshop 或专业排版软件里打开，不能拿来做站内预览。"""
    image_bytes = await image.read()
    try:
        img = PILImage.open(io.BytesIO(image_bytes)).convert("RGB")
    except Exception:
        raise HTTPException(status_code=400, detail="无法解析上传的图片")
    cmyk_img = img.convert("CMYK")
    buf = io.BytesIO()
    cmyk_img.save(buf, format="JPEG", quality=95)
    b64 = base64.b64encode(buf.getvalue()).decode()
    return {"image": f"data:image/jpeg;base64,{b64}"}


# "素材插画/文字生成"专用：只从参考图里框选出的一小块区域提炼风格类别（跟 reference-to-background
# 同一条纪律——只说类别/手法，不描述精确外形/不抄录原文文字），插画和文字两种分开写 instruction
# 是因为要防的具体表达完全不同：插画防的是"具体外形/姿势"，文字防的是"照抄原文这几个字"。
_ASSET_ILLUSTRATION_STYLE_INSTRUCTION = """You are extracting a STYLE-DESCRIPTION prompt for ONE isolated illustration/icon element cropped from a reference image, to be used as input for a text-to-image sticker generator. Describe ONLY the element's CATEGORY/TYPE and art style (e.g. "a paint palette with colorful paint blobs, flat cartoon illustration style, thick outline") as a mood-board brief would — not exact pose, not exact proportions, not exact color values. Output ONE short phrase in Chinese, no markdown, no quotation marks, no extra commentary."""

_ASSET_TEXT_STYLE_INSTRUCTION = """You are extracting the ARTISTIC RENDERING STYLE of a piece of text cropped from a reference image — its color scheme, outline/shadow/glow treatment, and letter-shape character (e.g. "彩虹渐变泡泡字，白色粗描边，卡通圆润造型") — for use as a style instruction for a text-to-image generator that will render COMPLETELY DIFFERENT words in this style. Do NOT mention, transcribe, or hint at the actual words/characters shown in the image. Output ONE short phrase in Chinese, no markdown, no quotation marks, no extra commentary."""


@router.post("/design/reference-to-asset")
async def design_reference_to_asset(
    image: UploadFile = File(...),
    region_x: float = Form(...),
    region_y: float = Form(...),
    region_width: float = Form(...),
    region_height: float = Form(...),
    asset_type: str = Form(...),
    text: str | None = Form(None),
    user: models.User = Depends(auth.get_current_user),
    db: Session = Depends(get_db),
):
    """参考图里框选一小块区域（插画元素或一段造型文字）→ 提炼风格类别 → 生成一张全新的独立
    素材 → 本地 rembg 抠成透明背景 → 自动存进素材库，返回可直接拖进画布的 PNG。
    跟 reference-to-background 是同一套边界纪律的另一个应用场景：插画防"抄具体外形"，
    文字防"抄原文文字"——用户自己填 text 参数，生成的是新内容，只是照抄了"手法"。"""
    if asset_type not in ("illustration", "text"):
        raise HTTPException(status_code=400, detail="asset_type 必须是 illustration 或 text")
    if asset_type == "text" and not (text and text.strip()):
        raise HTTPException(status_code=400, detail="文字类型需要提供 text 内容")
    _require_openlux()
    if asset_type == "text":
        await _moderate_text(text or "")

    image_bytes = await image.read()
    await _check_not_sensitive_document(image_bytes, image.content_type or "image/png", "素材/文字生成")
    try:
        img = PILImage.open(io.BytesIO(image_bytes)).convert("RGB")
    except Exception:
        raise HTTPException(status_code=400, detail="无法解析上传的图片")
    w, h = img.size
    box = (
        max(0, min(w, round(region_x * w))),
        max(0, min(h, round(region_y * h))),
        max(0, min(w, round((region_x + region_width) * w))),
        max(0, min(h, round((region_y + region_height) * h))),
    )
    if box[2] <= box[0] or box[3] <= box[1]:
        raise HTTPException(status_code=400, detail="选区无效")
    crop_buf = io.BytesIO()
    img.crop(box).save(crop_buf, format="PNG")
    crop_b64 = base64.b64encode(crop_buf.getvalue()).decode()

    ticket = billing.consume(db, user, "素材生成")
    try:
        return await _do_reference_to_asset(db, user, asset_type, text, crop_b64)
    except Exception:
        billing.refund_ticket(db, user, ticket)
        raise


async def _do_reference_to_asset(db, user, asset_type: str, text: str | None, crop_b64: str):
    instruction = _ASSET_ILLUSTRATION_STYLE_INSTRUCTION if asset_type == "illustration" else _ASSET_TEXT_STYLE_INSTRUCTION
    vision_res = await _post_openlux(
        f"{OPENLUX_BASE_URL}/chat/completions",
        timeout=60,
        headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
        json={
            "model": "gemini-3-flash-preview",
            "messages": [
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": instruction},
                        {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{crop_b64}"}},
                    ],
                }
            ],
        },
    )
    if vision_res.status_code >= 400:
        raise HTTPException(status_code=502, detail=f"素材风格分析失败：{vision_res.status_code} {vision_res.text}")
    style_description = vision_res.json()["choices"][0]["message"]["content"].strip()

    if asset_type == "illustration":
        gen_prompt = f"{style_description}，纯白色背景，不要文字，不要水印，贴纸风格，孤立单个物体，居中构图"
    else:
        gen_prompt = f'文字内容"{text.strip()}"，渲染风格：{style_description}，纯白色背景，不要多余装饰，孤立文字图形，居中构图'

    gen_res = await _post_openlux(
        f"{OPENLUX_BASE_URL}/images/generations",
        timeout=170,
        headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
        json={"model": "gpt-image-2", "prompt": gen_prompt, "n": 1, "size": "1024x1024"},
    )
    if gen_res.status_code >= 400:
        raise HTTPException(status_code=502, detail=f"素材生成失败：{gen_res.status_code} {gen_res.text}")
    gen_bytes = await _extract_openai_image_bytes(gen_res.json(), "素材生成")

    async with _HEAVY_MODEL_SEM:
        try:
            # 贴纸/文字图形是白底孤立单体，要干净的硬边，走 hard 档
            proc = await asyncio.to_thread(_run_bg_removal_subprocess, gen_bytes, "hard")
        except subprocess.TimeoutExpired:
            raise HTTPException(status_code=504, detail="抠图处理超时，请重试")
    if proc.returncode != 0:
        raise HTTPException(status_code=502, detail=f"抠图处理失败：{proc.stderr.decode(errors='ignore')[:500]}")

    category = "sticker-illustration" if asset_type == "illustration" else "sticker-text"
    asset = _persist_asset_bytes(db, user.id, category, proc.stdout)

    return {"assetId": asset.id, "url": f"/api/ai/generated/{asset.file_name}", "styleDescription": style_description}


def _flatten_chat_text(messages: list) -> str:
    """把 ChatMessage.content（可能是纯字符串，也可能是多模态 content-parts 数组）拍平成
    一段纯文本给合规分类器看——分类器只关心"这段文字想让 AI 写什么"，不需要图片部分。"""
    parts: list[str] = []
    for m in messages:
        content = m.content
        if isinstance(content, str):
            parts.append(content)
        elif isinstance(content, list):
            for item in content:
                if isinstance(item, dict) and item.get("type") == "text":
                    parts.append(str(item.get("text", "")))
    return "\n".join(parts)


@router.post("/chat/completions")
async def chat_completions(
    payload: ChatCompletionRequest,
    _user: models.User = Depends(auth.get_current_user),
):
    _require_openlux()
    await _moderate_text(_flatten_chat_text(payload.messages))
    res = await _post_openlux(
        f"{OPENLUX_BASE_URL}/chat/completions",
        timeout=60,
        headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
        json=payload.model_dump(exclude_none=True),
    )
    if res.status_code >= 400:
        raise HTTPException(status_code=502, detail=f"{res.status_code} {res.text}")
    return res.json()


@router.post("/video/generate")
async def video_generate(
    payload: VideoGenerateRequest,
    _user: models.User = Depends(auth.get_current_user),
):
    _require_vidu()
    async with httpx.AsyncClient(timeout=60) as client:
        res = await client.post(
            f"{VIDU_BASE_URL}/ent/v2/text2video",
            headers={"Authorization": f"Bearer {VIDU_API_KEY}", "Accept": "application/json"},
            json=payload.model_dump(),
        )
    if res.status_code >= 400:
        raise HTTPException(status_code=502, detail=f"视频生成接口请求失败：{res.status_code} {res.text}")
    return res.json()


@router.get("/video/tasks/{task_id}")
async def video_task_status(
    task_id: str,
    _user: models.User = Depends(auth.get_current_user),
):
    _require_vidu()
    async with httpx.AsyncClient(timeout=30) as client:
        res = await client.get(
            f"{VIDU_BASE_URL}/ent/v2/tasks/{task_id}/creations",
            headers={"Authorization": f"Bearer {VIDU_API_KEY}", "Accept": "application/json"},
        )
    if res.status_code >= 400:
        raise HTTPException(status_code=502, detail=f"视频任务状态查询失败：{res.status_code} {res.text}")
    return res.json()


def _clamp_num(value, lo: float, hi: float, default: float) -> float:
    try:
        num = float(value)
    except (TypeError, ValueError):
        return default
    return max(lo, min(hi, num))


def _nearest_image_size(width: float, height: float) -> str:
    target_ratio = (width / height) if height else 1.0
    _, w, h = min(_ASPECT_SIZES, key=lambda item: abs((item[1] / item[2]) - target_ratio))
    return f"{w}x{h}"


_ICON_LIST_SHAPES = ("circle", "square", "diamond")
_ICON_LIST_MAX_ITEMS = 12

# design/generate 和 design/layout 两个接口共用的元素类型/组件 JSON 结构说明——两边的
# system prompt 其余部分差别很大（一个是"帮我编内容"，一个是"内容都写好了，只排版"），
# 但"可以用哪些元素类型、每种长什么样"这部分是同一套渲染管线，两边必须保持一致，
# 抽出来是为了避免以后改了一边、另一边忘了同步改。
_COMPONENT_JSON_SCHEMA_DOC = (
    "可以使用四种元素类型：\n"
    "- text：{type:'text', x, y, width, text, fontSize, fontWeight, color, align, fontFamily}\n"
    "- image：{type:'image', x, y, width, height, imagePrompt}（用 imagePrompt 描述这张图应该是什么内容，不要写 src）\n"
    "- rect：{type:'rect', x, y, width, height, fill, rx}\n"
    "- group：结构化组件，两种 componentKind 可选，内容类型和它们匹配时优先用这个而不是堆多个 text：\n"
    "  1. icon-list（编号/勾选/要点清单，3~9 条短句时用这个）：\n"
    "     {type:'group', x, y, children:[], componentKind:'icon-list', componentData:[\n"
    "       {shape:'circle'|'square'|'diamond', color:'#c8161d', icon:'1', label:'一句话要点（12字以内）'}, ...\n"
    "     ]}\n"
    "     componentData 是数组，每项一条列表内容；icon 通常是序号（'1'/'2'）或单字（'✓'/'✕'），最多2个字符；\n"
    "     渲染时每项固定占 34px 高、180px 宽的文字区，label 太长会被截断，务必控制在 12 个汉字以内。\n"
    "  2. ribbon-title（分区小标题，通栏色块+居中白字）：\n"
    "     {type:'group', x, y, children:[], componentKind:'ribbon-title', componentData:[{text:'小标题（8字以内）', color:'#c8161d', width:220}]}\n"
    "     componentData 必须是单元素数组（哪怕只有一项也要包成数组），不能是裸对象。\n"
    "  什么时候该用 group 而不是普通 text：一段内容如果是"
    "「3条以上并列的短要点」，用 icon-list；一段内容如果是「引出下面一块信息的小标题」，用 ribbon-title；"
    "长段落说明文字、单独一句话的大标题，仍然用普通 text，不要什么都往组件里塞。\n"
)


def _sanitize_icon_list_data(raw: object) -> list[dict] | None:
    """对应前端 IconListDatum[]（CanvasStage.vue 的 buildIconList）——数组形状是硬要求，
    applyFieldEdit 双击编辑回写只认数组，裸对象会导致编辑框弹得出来但提交静默不生效。"""
    if not isinstance(raw, list) or not raw:
        return None
    items = []
    for it in raw[:_ICON_LIST_MAX_ITEMS]:
        if not isinstance(it, dict):
            continue
        label = str(it.get("label") or "").strip()[:24]
        if not label:
            continue
        shape = it.get("shape") if it.get("shape") in _ICON_LIST_SHAPES else "circle"
        items.append(
            {
                "shape": shape,
                "color": str(it.get("color") or "#c8161d"),
                "icon": str(it.get("icon") or "").strip()[:2] or "•",
                "label": label,
            }
        )
    return items or None


def _sanitize_ribbon_title_data(raw: object) -> list[dict] | None:
    """对应前端 RibbonTitleDatum[]（buildRibbonTitle）——同样必须是单元素数组，不是裸对象。"""
    if not isinstance(raw, list) or not raw or not isinstance(raw[0], dict):
        return None
    text = str(raw[0].get("text") or "").strip()[:20]
    if not text:
        return None
    item: dict = {"text": text, "color": str(raw[0].get("color") or "#c8161d")}
    width = raw[0].get("width")
    if isinstance(width, (int, float)):
        item["width"] = int(_clamp_num(width, 80, 1000, 220))
    return [item]


_TEXT_MAX_LINES = 8  # 单个文字元素允许撑到的最大行数，超过就截断——防止极端情况下的连锁下推把画布挤爆


def _truncate_to_max_lines(text: str, width: float, font_size: float) -> str:
    """兜底：AI 给的文字量跟它给的框完全不匹配（比如一大段话塞进一个小框）时，
    与其让它在渲染时无限撑高、把下面所有元素连环顶飞，不如硬截断加省略号。
    正常情况下走的是 _reflow_avoid_overlap 顺势下推，不会碰到这个分支。"""
    if estimate_text_lines(text, width, font_size) <= _TEXT_MAX_LINES:
        return text
    max_chars = max(1, chars_per_line(width, font_size) * _TEXT_MAX_LINES - 1)
    return text[:max_chars].rstrip() + "…"


def _element_width(el: dict) -> float:
    if el["type"] in ("text", "rect", "image"):
        return el["width"]
    if el["type"] == "group":
        if el.get("componentKind") == "ribbon-title":
            data = el.get("componentData") or [{}]
            return data[0].get("width", COMPONENT_SIZE["ribbon"]["defaultW"])
        if el.get("componentKind") == "icon-list":
            return COMPONENT_SIZE["iconList"]["badge"] + 10 + COMPONENT_SIZE["iconList"]["labelW"]
    return 200


def _ranges_overlap(a0: float, a1: float, b0: float, b1: float) -> bool:
    return a0 < b1 and b0 < a1


def _reflow_avoid_overlap(elements: list[dict]) -> None:
    """原地调整：AI 给文字元素的坐标是"凭感觉"定的，没有真的算过换行后会撑多高。
    Fabric 的 Textbox 不会裁切溢出内容，只会把自己撑高，直接压到下一个元素身上。
    这里按估算高度，把跟它横向范围有重叠（真正会被压到的，不是并排的另一栏）的
    后续元素顺势下移，避免遮挡/重叠。按数组顺序单趟扫描，靠"后面元素用的是已经
    调整过的 y"自然实现连锁下推，不需要额外的多轮迭代。"""
    for i, el in enumerate(elements):
        if el["type"] != "text":
            continue
        est_bottom = el["y"] + estimate_text_height(el["text"], el["width"], el["fontSize"])
        x0, x1 = el["x"], el["x"] + el["width"]
        for other in elements[i + 1 :]:
            ox0, ox1 = other["x"], other["x"] + _element_width(other)
            if not _ranges_overlap(x0, x1, ox0, ox1):
                continue
            if other["y"] < est_bottom:
                other["y"] = int(est_bottom) + 4


def _sanitize_design_elements(
    raw_elements: object,
    canvas_width: int,
    canvas_height: int,
    allowed_fonts: list[str],
    max_elements: int = _DESIGN_MAX_ELEMENTS,
) -> list[dict]:
    if not isinstance(raw_elements, list):
        return []
    sanitized: list[dict] = []
    for el in raw_elements[:max_elements]:
        if not isinstance(el, dict) or el.get("type") not in ("text", "image", "rect", "group"):
            continue
        x = int(_clamp_num(el.get("x"), 0, canvas_width, 0))
        y = int(_clamp_num(el.get("y"), 0, canvas_height, 0))
        width = int(_clamp_num(el.get("width"), 1, canvas_width, min(200, canvas_width)))

        if el["type"] == "text":
            text = str(el.get("text") or "").strip()[:200]
            if not text:
                continue
            font_family = el.get("fontFamily")
            if not isinstance(font_family, str) or font_family not in allowed_fonts:
                font_family = allowed_fonts[0] if allowed_fonts else None
            item = {
                "type": "text",
                "x": x,
                "y": y,
                "width": width,
                "text": text,
                "fontSize": int(_clamp_num(el.get("fontSize"), 10, 160, 32)),
                "fontWeight": str(el.get("fontWeight") or "normal"),
                "color": str(el.get("color") or "#000000"),
                "align": el.get("align") if el.get("align") in ("left", "center", "right") else "left",
            }
            if font_family:
                item["fontFamily"] = font_family
            item["text"] = _truncate_to_max_lines(item["text"], item["width"], item["fontSize"])
            sanitized.append(item)
        elif el["type"] == "rect":
            height = int(_clamp_num(el.get("height"), 1, canvas_height, min(200, canvas_height)))
            sanitized.append(
                {
                    "type": "rect",
                    "x": x,
                    "y": y,
                    "width": width,
                    "height": height,
                    "fill": str(el.get("fill") or "#000000"),
                    "rx": int(_clamp_num(el.get("rx"), 0, 200, 0)),
                }
            )
        elif el["type"] == "image":
            height = int(_clamp_num(el.get("height"), 1, canvas_height, min(200, canvas_height)))
            image_prompt = el.get("imagePrompt")
            if not isinstance(image_prompt, str) or not image_prompt.strip():
                continue
            sanitized.append(
                {
                    "type": "image",
                    "x": x,
                    "y": y,
                    "width": width,
                    "height": height,
                    "imagePrompt": image_prompt.strip()[:300],
                }
            )
        elif el["type"] == "group":
            kind = el.get("componentKind")
            if kind == "icon-list":
                data = _sanitize_icon_list_data(el.get("componentData"))
            elif kind == "ribbon-title":
                data = _sanitize_ribbon_title_data(el.get("componentData"))
            else:
                continue
            if data is None:
                continue
            sanitized.append({"type": "group", "x": x, "y": y, "children": [], "componentKind": kind, "componentData": data})
    _reflow_avoid_overlap(sanitized)
    return sanitized


@router.post("/design/generate")
async def design_generate(
    payload: DesignGenerateRequest,
    _user: models.User = Depends(auth.get_current_user),
):
    """
    AI 一键生成设计：先让文字模型拟一版版式 JSON（图片元素只给 imagePrompt），
    再并发调图片生成把 imagePrompt 换成真实 src，最后拼装成 CanvasElement[] 返回。
    未经真实联调验证——如果调用报错或解析失败，把报错信息发给我，按实际返回结构调整。
    """
    _require_openlux()

    allowed_fonts = [f.value for f in payload.fonts] or ["sans-serif"]
    font_list_text = "\n".join(f"- {f.label}：{f.value}" for f in payload.fonts) or "- 默认：sans-serif"

    system_prompt = (
        f"你是专业的平面设计师。请为一个宽 {payload.canvas_width}px、高 {payload.canvas_height}px 的画布设计一版海报/宣传图版式。\n"
        + _COMPONENT_JSON_SCHEMA_DOC
        + f"可选字体（fontFamily 必须从下面列表里原样选一个，不要自己编）：\n{font_list_text}\n\n"
        "只返回一个严格的 JSON 对象，不要 markdown 代码块，不要任何多余说明文字，形如：\n"
        '{"background": "#ffffff", "elements": [ ... ]}\n'
        "要求：elements 数量 4~12 个（一个 group 算一个元素，不受它内部 componentData 条数影响）；"
        "所有坐标、宽高不能超出画布范围；文案要贴合用户描述的主题，"
        "写真实、具体的中文文案，不要用「标题」「正文」这类占位字样；整体版面要有主次层次（标题/副标题/正文/装饰色块）。\n"
        "字体选择要有主次区分，不要所有文字都用同一个字体：标题、副标题这类需要吸引注意力的文字，"
        "优先挑选列表里风格醒目/有辨识度的字体（比如手写体、圆体、书法风格），不要默认全用「默认」这个选项；"
        "正文说明、联系方式这类信息性文字，选清晰易读的字体即可。同一版设计里可以搭配 2~3 种不同字体，避免单调。"
    )

    async with httpx.AsyncClient(timeout=60) as client:
        chat_res = await client.post(
            f"{OPENLUX_BASE_URL}/chat/completions",
            headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
            json={
                "model": "gemini-3-flash-preview",
                "messages": [
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": payload.prompt},
                ],
            },
        )
    if chat_res.status_code >= 400:
        raise HTTPException(status_code=502, detail=f"设计生成接口请求失败：{chat_res.status_code} {chat_res.text}")

    content = chat_res.json().get("choices", [{}])[0].get("message", {}).get("content", "")
    match = re.search(r"\{[\s\S]*\}", content)
    try:
        parsed = json.loads(match.group(0) if match else content)
    except (json.JSONDecodeError, AttributeError):
        raise HTTPException(status_code=502, detail="AI 设计生成解析失败，请重试")

    if not isinstance(parsed, dict):
        raise HTTPException(status_code=502, detail="AI 设计生成解析失败，请重试")

    background = str(parsed.get("background") or "#ffffff")
    elements = _sanitize_design_elements(
        parsed.get("elements"), payload.canvas_width, payload.canvas_height, allowed_fonts
    )

    async def _fill_image(el: dict) -> dict | None:
        size = _nearest_image_size(el["width"], el["height"])
        try:
            async with httpx.AsyncClient(timeout=120) as client:
                res = await client.post(
                    f"{OPENLUX_BASE_URL}/images/generations",
                    headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
                    json={"model": "gpt-image-2", "prompt": el["imagePrompt"], "n": 1, "size": size},
                )
            if res.status_code >= 400:
                return None
            data = (res.json().get("data") or [None])[0]
            if not data:
                return None
            src = data.get("url") or (f"data:image/png;base64,{data['b64_json']}" if data.get("b64_json") else None)
            if not src:
                return None
            return {"type": "image", "x": el["x"], "y": el["y"], "width": el["width"], "height": el["height"], "src": src}
        except httpx.HTTPError:
            return None

    image_indices = [i for i, el in enumerate(elements) if el["type"] == "image"]
    filled = (
        await asyncio.gather(*(_fill_image(elements[i]) for i in image_indices)) if image_indices else []
    )
    filled_by_index = dict(zip(image_indices, filled))

    final_elements = []
    for i, el in enumerate(elements):
        if el["type"] == "image":
            replacement = filled_by_index.get(i)
            if replacement is not None:
                final_elements.append(replacement)
            # 图片生成失败就整个丢掉这个槽位，不让整个请求失败
        else:
            final_elements.append(el)

    return {"background": background, "elements": final_elements}


def _build_layout_prompt(canvas_width: int, canvas_height: int, blocks: list[str] | None, raw_text: str | None) -> tuple[str, str]:
    system_prompt = (
        "你是专业的排版设计师。下面是用户已经写好的正文内容——这是最重要的规则：\n"
        "你绝对不能改写、删减、替换、精简、润色、翻译用户提供的原文任何一个字，你的任务只有排版，不是写作。\n\n"
        f"具体要做两件事，画布宽 {canvas_width}px、高 {canvas_height}px：\n"
        "1. 判断内容在语义上怎么拆成几个逻辑段落（如果输入已经按段落给出，直接按给定的分段来，不要重新合并/拆分）。"
        "每段归类成以下几种之一：「标题」（一句话，通常最短最醒目）、「要点列表」（3条以上并列的短句/短语）、"
        "「正文说明」（一段连续叙述）。\n"
        "2. 给每个逻辑段落分配版面位置和展示方式：\n"
        + _COMPONENT_JSON_SCHEMA_DOC
        + "这个模式暂不支持自动配图——不要生成 type:'image' 的元素，只用 text/rect/group 这三种。\n"
        + "\n再次强调最重要的规则：\n"
        "- 每个元素的 text/label 字段的内容，必须是从用户原文里逐字截取的连续一段，不能是你概括/重写/翻译/精简后的版本，一个字都不能改。\n"
        "- 唯一允许的操作是“在原文的自然断句处切开”（按标点、按换行、按逻辑段落切分成多段），切分内部不能改字、加字、减字。\n"
        "- 如果某一段原文太长装不进一个元素，可以原样拆成前后衔接的多个元素，但拼起来必须跟原文一字不差，不能省略中间内容、不能概括。\n"
        "- 不需要凭空加装饰性文案（比如联系方式、口号）——原文没有的内容，版面上也不应该出现。\n"
        "- 内容完整性优先于版面美观：宁可少加几个纯装饰用的色块/配图，也不能因为元素数量超限把用户原文的某一段挤丢。\n"
        "只返回一个严格的 JSON 对象，不要 markdown 代码块，不要任何多余说明文字，形如：\n"
        '{"background": "#ffffff", "elements": [ ... ]}\n'
    )
    if blocks:
        content_desc = "\n".join(f"[{i}] {b}" for i, b in enumerate(blocks))
        user_prompt = f"用户已经把内容分好段了，请分别给每段选合适的展示方式和位置，不要合并/拆分/改写这些段落原文：\n\n{content_desc}"
    else:
        user_prompt = f"用户原文（还没有分段，请你先按语义切分，再分别排版，原文一个字都不能改）：\n\n{raw_text}"
    return system_prompt, user_prompt


def _check_verbatim_fidelity(elements: list[dict], source_text: str) -> dict:
    """核对排版结果里每一段文字，是不是真的从用户原文逐字截取的，不是模型概括/改写过的版本。
    不能只在 prompt 里嘱咐一句就信——这条经验是从联网搜索那条链路的 quote 校验机制照搬过来的：
    模型自己说"我没有改写"不可信，字符串层面能校验的东西就不要指望模型自觉。
    允许兜底截断留下的结尾"…"不算破坏逐字性（那是 _sanitize_design_elements 自己加的安全网，不是模型编的）。"""
    normalized_source = content_research.normalize_for_match(source_text)
    checks: list[dict] = []

    def check_one(kind: str, text: str) -> None:
        candidate = text[:-1] if text.endswith("…") else text
        verbatim = bool(candidate) and content_research.normalize_for_match(candidate) in normalized_source
        checks.append({"type": kind, "text": text, "verbatim": verbatim})

    for el in elements:
        if el["type"] == "text":
            check_one("text", el["text"])
        elif el["type"] == "group":
            if el.get("componentKind") == "ribbon-title":
                for item in el.get("componentData") or []:
                    check_one("ribbon-title", str(item.get("text", "")))
            elif el.get("componentKind") == "icon-list":
                for item in el.get("componentData") or []:
                    check_one("icon-list", str(item.get("label", "")))

    verbatim_count = sum(1 for c in checks if c["verbatim"])
    return {
        "total": len(checks),
        "verbatim_count": verbatim_count,
        "all_verbatim": verbatim_count == len(checks) if checks else True,
        "checks": checks,
    }


@router.post("/design/layout")
async def design_layout(
    payload: DesignLayoutRequest,
    _user: models.User = Depends(auth.get_current_user),
):
    """跟 /design/generate 是两条独立的链路，不共用 system prompt——那边是"帮我编内容"，
    这边是"内容都写好了，只排版、只选组件"。逐字校验见 _check_verbatim_fidelity。"""
    _require_openlux()

    blocks = [b.strip() for b in (payload.blocks or []) if b and b.strip()]
    raw_text = (payload.raw_text or "").strip()
    if not blocks and not raw_text:
        raise HTTPException(status_code=400, detail="raw_text 和 blocks 至少要提供一个")
    source_text = "".join(blocks) if blocks else raw_text

    system_prompt, user_prompt = _build_layout_prompt(
        payload.canvas_width, payload.canvas_height, blocks or None, raw_text or None
    )

    async with httpx.AsyncClient(timeout=90) as client:
        chat_res = await client.post(
            f"{OPENLUX_BASE_URL}/chat/completions",
            headers={"Authorization": f"Bearer {OPENLUX_API_KEY}"},
            json={
                "model": "gemini-3-flash-preview",
                "messages": [
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_prompt},
                ],
            },
        )
    if chat_res.status_code >= 400:
        raise HTTPException(status_code=502, detail=f"排版接口请求失败：{chat_res.status_code} {chat_res.text}")

    content = chat_res.json().get("choices", [{}])[0].get("message", {}).get("content", "")
    match = re.search(r"\{[\s\S]*\}", content)
    try:
        parsed = json.loads(match.group(0) if match else content)
    except (json.JSONDecodeError, AttributeError):
        raise HTTPException(status_code=502, detail="排版结果解析失败，请重试")
    if not isinstance(parsed, dict):
        raise HTTPException(status_code=502, detail="排版结果解析失败，请重试")

    allowed_fonts = [f.value for f in payload.fonts] or ["sans-serif"]
    background = str(parsed.get("background") or "#ffffff")
    elements = _sanitize_design_elements(
        parsed.get("elements"), payload.canvas_width, payload.canvas_height, allowed_fonts, max_elements=_LAYOUT_MAX_ELEMENTS
    )
    # 双保险：prompt 里已经说了这个模式不生成 image 元素，但指令不可靠（这次会话已经在别的地方
    # 踩过这个坑），万一模型还是给了，这里直接过滤掉——总比留一个没有 src 的残缺元素在版面里好。
    elements = [el for el in elements if el["type"] != "image"]
    fidelity = _check_verbatim_fidelity(elements, source_text)

    return {"background": background, "elements": elements, "fidelity": fidelity}


@router.post("/design/layout-preset")
async def design_layout_preset(
    payload: LayoutPresetRequest,
    _user: models.User = Depends(auth.get_current_user),
):
    """参数化排版预设——不调用任何 AI，纯确定性代码（见 app/layout_presets.py），
    输入已经是分好类的结构化内容，系统只管按选定结构算坐标。跟 /design/layout 的区别：
    这里没有 AI 猜坐标那种"数值判断不准"的风险，也没有网络延迟。"""
    if payload.structure == "bullet-list":
        if not payload.items:
            raise HTTPException(status_code=400, detail="bullet-list 结构需要提供 items")
        result = layout_presets.build_bullet_list(payload.canvas_width, payload.canvas_height, payload.title, payload.intro, payload.items)
    elif payload.structure == "dense-board":
        if not payload.sections:
            raise HTTPException(status_code=400, detail="dense-board 结构需要提供 sections")
        sections = [{"heading": s.heading, "items": s.items} for s in payload.sections]
        result = layout_presets.build_dense_board(
            payload.canvas_width,
            payload.canvas_height,
            payload.title,
            sections,
            include_title=payload.include_title,
            top_offset=payload.top_offset,
            colors=payload.colors,
        )
    else:
        raise HTTPException(status_code=400, detail=f"未知的 structure：{payload.structure}")
    return result


@router.post("/content/research")
async def content_research_endpoint(
    payload: ContentResearchRequest,
    _user: models.User = Depends(auth.get_current_user),
):
    """联网搜索一个主题 -> 抓取排名靠前页面的全文 -> 喂给 gemini-3-flash-preview 提炼成
    {claim, quote, source_url, site_name, confidence} 列表。confidence 不是模型自报的分数，
    是拿 quote 去抓到的原文里做字符串校验算出来的（verified/extracted_unverified/not_found，
    具体规则见 content_research.py 顶部注释）。"""
    topic = payload.topic.strip()
    if not topic:
        raise HTTPException(status_code=400, detail="topic 不能为空")
    try:
        return await content_research.research_topic(topic)
    except content_research.ResearchUnavailable as exc:
        raise HTTPException(status_code=503, detail=str(exc))
    except RuntimeError as exc:
        raise HTTPException(status_code=502, detail=str(exc))
